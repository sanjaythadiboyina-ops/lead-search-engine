import os

import csv

import re

import base64

import hashlib

import hmac

import json

import secrets

from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from io import BytesIO

from pathlib import Path

from typing import Optional, List

from collections import defaultdict, deque

from urllib.parse import quote_plus, urlparse

from urllib import robotparser
from html.parser import HTMLParser

import jwt

import psycopg

import requests

from dotenv import load_dotenv

from fastapi import FastAPI, HTTPException, Depends, Header, Query

from fastapi.middleware.cors import CORSMiddleware

from fastapi.responses import FileResponse, StreamingResponse, Response

from openpyxl import Workbook

from pydantic import BaseModel, Field

# ============================================================

# Configuration

# ============================================================

BASE_DIR = Path(__file__).resolve().parent

PROJECT_DIR = BASE_DIR.parent

FRONTEND_FILE = PROJECT_DIR / "frontend" / "index.html"

load_dotenv(PROJECT_DIR / ".env")

load_dotenv()

DATABASE_URL = os.getenv("DATABASE_URL")

JWT_SECRET = os.getenv("JWT_SECRET") or secrets.token_urlsafe(32)

JWT_ALGORITHM = os.getenv("JWT_ALGORITHM", "HS256")

LATLNG_API_KEY = os.getenv("LATLNG_API_KEY")

SERPAPI_API_KEY = os.getenv("SERPAPI_API_KEY")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
FIREBASE_API_KEY = os.getenv("FIREBASE_API_KEY")

JWT_EXPIRE_HOURS = int(os.getenv("JWT_EXPIRE_HOURS", "168"))

LATLNG_GEOCODE_URL = "https://api.latlng.work/api"

LATLNG_PLACES_NEARBY_URL = "https://api.latlng.work/v1/places/nearby"

LATLNG_PLACES_SEARCH_URL = "https://api.latlng.work/v1/places/search"

GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")

SEARCH_RADIUS_METERS = min(5000, max(250, int(os.getenv("SEARCH_RADIUS_METERS", "5000"))))

MAX_LIMIT = min(100, max(1, int(os.getenv("MAX_SEARCH_RESULTS", "50"))))

CITY_WIDE_MAX_ZONES = min(9, max(1, int(os.getenv("CITY_WIDE_MAX_ZONES", "9"))))

CITY_WIDE_GRID_SPACING_METERS = min(4500, max(1000, int(os.getenv("CITY_WIDE_GRID_SPACING_METERS", "3500"))))

ENRICH_TIMEOUT = min(15, max(3, int(os.getenv("ENRICH_TIMEOUT_SECONDS", "8"))))

app = FastAPI(title="Lead Searching Engine", version="4.0")
_SEARCH_EVENTS = defaultdict(deque)
SEARCH_WINDOW_SECONDS = 300
SEARCH_MAX_REQUESTS = 20

app.add_middleware(

    CORSMiddleware,

    allow_origins=["*"],

    allow_credentials=True,

    allow_methods=["*"],

    allow_headers=["*"],

)

CATEGORY_MAP = {

    "Restaurant": "restaurant", "Cafe": "cafe", "Bakery": "bakery", "Bar": "bar",

    "Gym": "fitness_centre", "Salon": "beauty_salon", "Dentist": "dentist", "Clinic": "clinic",

    "Pharmacy": "pharmacy", "Hotel": "hotel", "Real Estate": "estate_agent", "Supermarket": "supermarket",

    "Jewellery": "jewellery", "Furniture": "furniture", "Car Repair": "car_repair",

    "Travel Agency": "travel_agency", "Photography": "photographer", "Event Management": "event_venue",

    "Interior Design": "interior_design", "Boutique": "clothes",

}

VALID_STATUSES = {"New", "Contacted", "Follow-up", "Negotiation", "Converted", "Lost"}

VALID_ROLES = {"Admin", "Agent"}


def is_admin(user: dict) -> bool:
    return str(user.get("role") or "Agent") == "Admin"


def require_admin(user: dict) -> None:
    if not is_admin(user):
        raise HTTPException(status_code=403, detail="Admin permission required.")


def is_team_leader(user: dict) -> bool:
    return False

def accessible_user_clause(user: dict, alias: str = "u"):
    prefix = f"{alias}." if alias else ""
    if is_admin(user):
        return "TRUE", []
    return f"{prefix}id=%s", [user["id"]]

def lead_access_clause(user: dict, alias: str = ""):
    prefix = f"{alias}." if alias else ""
    if is_admin(user):
        return "TRUE", []
    return f"{prefix}assigned_to=%s", [user["id"]]


def lead_owner_check(cur, lead_id: int, user: dict, for_update: bool = False):
    scope, params = lead_access_clause(user)
    sql = f"SELECT id,assigned_to,user_id FROM leads WHERE id=%s AND {scope}"
    if for_update:
        sql += " FOR UPDATE"
    cur.execute(sql, [lead_id, *params])
    row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Lead not found or you do not have access to it.")
    return row


def _notification_text(cur, user_id: int, lead_id: Optional[int], action: str, details: Optional[dict] = None):
    details = details or {}
    cur.execute("SELECT name,email,COALESCE(role,'Agent') FROM users WHERE id=%s", (user_id,))
    actor = cur.fetchone()
    actor_name = (actor[0] if actor and actor[0] else actor[1] if actor else None) or 'Agent'
    lead_name = None
    if lead_id is not None:
        cur.execute("SELECT name FROM leads WHERE id=%s", (lead_id,))
        row = cur.fetchone()
        lead_name = row[0] if row else None
    label = action.replace('_', ' ').strip().capitalize()
    title = f"{actor_name} · {label}"
    message = f"{actor_name} {label.lower()}" + (f" for {lead_name}" if lead_name else "") + "."
    if action in ('status_changed','bulk_status_changed') and details.get('status'):
        message = f"{actor_name} changed {lead_name or 'a lead'} status to {details['status']}."
    elif action == 'call_logged' and details.get('outcome'):
        message = f"{actor_name} logged a call for {lead_name or 'a lead'}: {details['outcome']}."
    elif action == 'followup_changed' and details.get('next_followup_date'):
        message = f"{actor_name} scheduled a follow-up for {lead_name or 'a lead'}."
    elif action == 'notes_changed':
        message = f"{actor_name} updated notes for {lead_name or 'a lead'}."
    elif action == 'lead_updated':
        message = f"{actor_name} updated {lead_name or 'a lead'}."
    return actor, title, message

def notify_admins_of_agent_activity(cur, user_id: int, lead_id: Optional[int], action: str, details: Optional[dict] = None):
    actor, title, message = _notification_text(cur, user_id, lead_id, action, details)
    if not actor or actor[2] != 'Agent': return
    cur.execute("SELECT id FROM users WHERE role='Admin' AND COALESCE(active,TRUE)=TRUE")
    for (admin_id,) in cur.fetchall():
        cur.execute("INSERT INTO notifications(recipient_id,actor_id,lead_id,notification_type,title,message) VALUES(%s,%s,%s,%s,%s,%s)",
                    (admin_id, user_id, lead_id, 'agent_activity', title, message))

def log_activity(cur, user_id: int, lead_id: Optional[int], action: str, details: Optional[dict] = None):
    details = details or {}
    cur.execute("INSERT INTO activity_logs(user_id,lead_id,action,details) VALUES(%s,%s,%s,%s::jsonb)",
                (user_id, lead_id, action, json.dumps(details, default=str)))
    notify_admins_of_agent_activity(cur, user_id, lead_id, action, details)


def whatsapp_url(phone: Optional[str]) -> Optional[str]:
    if not phone:
        return None
    digits = re.sub(r"\D", "", str(phone))
    if digits.startswith("00"):
        digits = digits[2:]
    if len(digits) == 10:
        digits = "91" + digits
    elif digits.startswith("0") and len(digits) == 11:
        digits = "91" + digits[1:]
    if len(digits) < 10:
        return None
    return f"https://wa.me/{digits}"


def user_timezone_name() -> str:
    return os.getenv("APP_TIMEZONE", "Asia/Kolkata")


def local_today_utc_bounds() -> tuple[datetime, datetime, str]:
    try:
        tz = ZoneInfo(user_timezone_name())
    except Exception:
        tz = timezone.utc
    now_local = datetime.now(tz)
    start_local = datetime.combine(now_local.date(), datetime.min.time(), tzinfo=tz)
    end_local = start_local + timedelta(days=1)
    return start_local.astimezone(timezone.utc), end_local.astimezone(timezone.utc), now_local.date().isoformat()

def local_date_range_utc(date_from: Optional[str] = None, date_to: Optional[str] = None):
    """Convert user-selected local calendar dates to UTC bounds using APP_TIMEZONE."""
    if not date_from and not date_to:
        return None, None
    try:
        tz = ZoneInfo(user_timezone_name())
    except Exception:
        tz = timezone.utc
    try:
        start_local = datetime.fromisoformat(date_from).replace(tzinfo=tz) if date_from else None
        end_local = (datetime.fromisoformat(date_to).replace(tzinfo=tz) + timedelta(days=1)) if date_to else None
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date filter. Use YYYY-MM-DD.")
    return (start_local.astimezone(timezone.utc) if start_local else None,
            end_local.astimezone(timezone.utc) if end_local else None)

def excel_safe_value(value):
    """Convert timestamps to APP_TIMEZONE before writing timezone-naive Excel datetimes."""
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            try:
                value = value.astimezone(ZoneInfo(user_timezone_name()))
            except Exception:
                pass
            return value.replace(tzinfo=None)
        return value
    return value

# ============================================================

# Frontend

# ============================================================

@app.get("/")

def home():

    if not FRONTEND_FILE.exists():

        raise HTTPException(status_code=500, detail="frontend/index.html was not found.")

    return FileResponse(FRONTEND_FILE)

# ============================================================

# Database / safe migrations

# ============================================================

def db_conn():

    if not DATABASE_URL:

        raise HTTPException(status_code=500, detail="DATABASE_URL is not configured.")

    return psycopg.connect(DATABASE_URL)



def _add_columns(cur, table: str, migrations: dict):

    existing = {

        row[0] for row in cur.execute(

            """SELECT column_name FROM information_schema.columns

               WHERE table_schema='public' AND table_name=%s""", (table,)

        ).fetchall()

    }

    for column, sql in migrations.items():

        if column not in existing:

            try:

                cur.execute(sql)

            except Exception as exc:

                print(f"Migration warning {table}.{column}: {exc}")



def init_db():

    if not DATABASE_URL:

        print("WARNING: DATABASE_URL is not configured.")

        return

    with db_conn() as conn:

        with conn.cursor() as cur:

            cur.execute("""CREATE TABLE IF NOT EXISTS users (

                id SERIAL PRIMARY KEY, email TEXT UNIQUE NOT NULL, name TEXT,

                password_hash TEXT NOT NULL, salt TEXT NOT NULL,

                created_at TIMESTAMPTZ DEFAULT NOW(), role TEXT DEFAULT 'Agent', team_leader_id INTEGER REFERENCES users(id) ON DELETE SET NULL, active BOOLEAN DEFAULT TRUE, firebase_uid TEXT, phone TEXT, approval_status TEXT DEFAULT 'Approved', auth_provider TEXT DEFAULT 'password')""")

            _add_columns(cur, "users", {

                "name": "ALTER TABLE users ADD COLUMN name TEXT",

                "password_hash": "ALTER TABLE users ADD COLUMN password_hash TEXT",

                "salt": "ALTER TABLE users ADD COLUMN salt TEXT",

                "created_at": "ALTER TABLE users ADD COLUMN created_at TIMESTAMPTZ DEFAULT NOW()",
                "role": "ALTER TABLE users ADD COLUMN role TEXT DEFAULT 'Agent'",
                "team_leader_id": "ALTER TABLE users ADD COLUMN team_leader_id INTEGER REFERENCES users(id) ON DELETE SET NULL",
                "active": "ALTER TABLE users ADD COLUMN active BOOLEAN DEFAULT TRUE",
                "firebase_uid": "ALTER TABLE users ADD COLUMN firebase_uid TEXT",
                "phone": "ALTER TABLE users ADD COLUMN phone TEXT",
                "approval_status": "ALTER TABLE users ADD COLUMN approval_status TEXT DEFAULT 'Approved'",
                "auth_provider": "ALTER TABLE users ADD COLUMN auth_provider TEXT DEFAULT 'password'",

            })

            cur.execute("""CREATE TABLE IF NOT EXISTS leads (

                id SERIAL PRIMARY KEY,

                user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,

                place_id TEXT,

                provider_place_id TEXT,

                name TEXT NOT NULL,

                address TEXT, city TEXT, region TEXT, country TEXT,

                latitude DOUBLE PRECISION, longitude DOUBLE PRECISION,

                phone TEXT, website TEXT, email TEXT,

                instagram TEXT, facebook TEXT, twitter TEXT, linkedin TEXT, youtube TEXT, tiktok TEXT,
                social_evidence JSONB DEFAULT '[]'::jsonb, social_status TEXT,

                category TEXT, rating DOUBLE PRECISION, reviews INTEGER,

                maps_url TEXT, source TEXT DEFAULT 'LatLng Places',

                sources JSONB DEFAULT '[]'::jsonb,

                lead_score INTEGER, lead_type TEXT,

                ai_recommendation TEXT, ai_reason TEXT,

                status TEXT DEFAULT 'New', notes TEXT,
                assigned_to INTEGER REFERENCES users(id) ON DELETE SET NULL,
                last_contacted_date TIMESTAMPTZ, next_followup_date TIMESTAMPTZ,
                is_duplicate BOOLEAN DEFAULT FALSE, duplicate_of INTEGER REFERENCES leads(id) ON DELETE SET NULL,

                enrichment_source TEXT, enrichment_at TIMESTAMPTZ,

                created_at TIMESTAMPTZ DEFAULT NOW(), updated_at TIMESTAMPTZ DEFAULT NOW())""")

            _add_columns(cur, "leads", {

                "user_id": "ALTER TABLE leads ADD COLUMN user_id INTEGER REFERENCES users(id) ON DELETE CASCADE",

                "place_id": "ALTER TABLE leads ADD COLUMN place_id TEXT",

                "provider_place_id": "ALTER TABLE leads ADD COLUMN provider_place_id TEXT",

                "city": "ALTER TABLE leads ADD COLUMN city TEXT",

                "region": "ALTER TABLE leads ADD COLUMN region TEXT",

                "country": "ALTER TABLE leads ADD COLUMN country TEXT",

                "latitude": "ALTER TABLE leads ADD COLUMN latitude DOUBLE PRECISION",

                "longitude": "ALTER TABLE leads ADD COLUMN longitude DOUBLE PRECISION",

                "email": "ALTER TABLE leads ADD COLUMN email TEXT",

                "instagram": "ALTER TABLE leads ADD COLUMN instagram TEXT",

                "facebook": "ALTER TABLE leads ADD COLUMN facebook TEXT",

                "twitter": "ALTER TABLE leads ADD COLUMN twitter TEXT",
                "linkedin": "ALTER TABLE leads ADD COLUMN linkedin TEXT",
                "youtube": "ALTER TABLE leads ADD COLUMN youtube TEXT",
                "tiktok": "ALTER TABLE leads ADD COLUMN tiktok TEXT",
                "social_evidence": "ALTER TABLE leads ADD COLUMN social_evidence JSONB DEFAULT '[]'::jsonb",
                "social_status": "ALTER TABLE leads ADD COLUMN social_status TEXT",

                "maps_url": "ALTER TABLE leads ADD COLUMN maps_url TEXT",

                "source": "ALTER TABLE leads ADD COLUMN source TEXT DEFAULT 'LatLng Places'",

                "sources": "ALTER TABLE leads ADD COLUMN sources JSONB DEFAULT '[]'::jsonb",

                "lead_score": "ALTER TABLE leads ADD COLUMN lead_score INTEGER",

                "lead_type": "ALTER TABLE leads ADD COLUMN lead_type TEXT",

                "ai_recommendation": "ALTER TABLE leads ADD COLUMN ai_recommendation TEXT",

                "ai_reason": "ALTER TABLE leads ADD COLUMN ai_reason TEXT",

                "status": "ALTER TABLE leads ADD COLUMN status TEXT DEFAULT 'New'",
                "assigned_to": "ALTER TABLE leads ADD COLUMN assigned_to INTEGER REFERENCES users(id) ON DELETE SET NULL",
                "last_contacted_date": "ALTER TABLE leads ADD COLUMN last_contacted_date TIMESTAMPTZ",
                "next_followup_date": "ALTER TABLE leads ADD COLUMN next_followup_date TIMESTAMPTZ",
                "is_duplicate": "ALTER TABLE leads ADD COLUMN is_duplicate BOOLEAN DEFAULT FALSE",
                "duplicate_of": "ALTER TABLE leads ADD COLUMN duplicate_of INTEGER REFERENCES leads(id) ON DELETE SET NULL",

                "notes": "ALTER TABLE leads ADD COLUMN notes TEXT",

                "enrichment_source": "ALTER TABLE leads ADD COLUMN enrichment_source TEXT",

                "enrichment_at": "ALTER TABLE leads ADD COLUMN enrichment_at TIMESTAMPTZ",
                "updated_at": "ALTER TABLE leads ADD COLUMN updated_at TIMESTAMPTZ DEFAULT NOW()",

            })

            # Preserve old data and old group memberships. Never drop group_leads on startup.

            cur.execute("""CREATE TABLE IF NOT EXISTS search_history (

                id SERIAL PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,

                city TEXT, category TEXT, keyword TEXT, city_wide BOOLEAN DEFAULT FALSE,

                result_count INTEGER, sources_used JSONB DEFAULT '[]'::jsonb, created_at TIMESTAMPTZ DEFAULT NOW(), role TEXT DEFAULT 'Agent')""")

            _add_columns(cur, "search_history", {

                "keyword": "ALTER TABLE search_history ADD COLUMN keyword TEXT",

                "city_wide": "ALTER TABLE search_history ADD COLUMN city_wide BOOLEAN DEFAULT FALSE",
                "sources_used": "ALTER TABLE search_history ADD COLUMN sources_used JSONB DEFAULT '[]'::jsonb",

            })

            cur.execute("""CREATE TABLE IF NOT EXISTS lead_groups (

                id SERIAL PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,

                name TEXT NOT NULL, created_at TIMESTAMPTZ DEFAULT NOW())""")

            cur.execute("""CREATE TABLE IF NOT EXISTS group_leads (

                id SERIAL PRIMARY KEY,

                group_id INTEGER NOT NULL REFERENCES lead_groups(id) ON DELETE CASCADE,

                lead_id INTEGER NOT NULL REFERENCES leads(id) ON DELETE CASCADE,

                created_at TIMESTAMPTZ DEFAULT NOW(), UNIQUE(group_id, lead_id))""")

            cur.execute("CREATE INDEX IF NOT EXISTS leads_user_created_idx ON leads(user_id, created_at DESC)")

            cur.execute("CREATE INDEX IF NOT EXISTS leads_user_city_idx ON leads(user_id, city)")

            cur.execute("CREATE INDEX IF NOT EXISTS leads_user_place_idx ON leads(user_id, place_id)")

            cur.execute("CREATE INDEX IF NOT EXISTS leads_assigned_status_idx ON leads(assigned_to,status,created_at DESC)")
            cur.execute("UPDATE leads SET status='New' WHERE status='Not Contacted'")
            cur.execute("UPDATE leads SET status='Negotiation' WHERE status='Interested'")
            cur.execute("""CREATE TABLE IF NOT EXISTS call_logs (
                id SERIAL PRIMARY KEY, lead_id INTEGER NOT NULL REFERENCES leads(id) ON DELETE CASCADE,
                called_by INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                call_date TIMESTAMPTZ DEFAULT NOW(), call_outcome TEXT NOT NULL, notes TEXT,
                created_at TIMESTAMPTZ DEFAULT NOW())""")
            cur.execute("""CREATE TABLE IF NOT EXISTS activity_logs (
                id SERIAL PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                lead_id INTEGER REFERENCES leads(id) ON DELETE CASCADE, action TEXT NOT NULL,
                details JSONB DEFAULT '{}'::jsonb, created_at TIMESTAMPTZ DEFAULT NOW())""")
            cur.execute("""CREATE TABLE IF NOT EXISTS notifications (
                id SERIAL PRIMARY KEY, recipient_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                actor_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                lead_id INTEGER REFERENCES leads(id) ON DELETE SET NULL,
                notification_type TEXT NOT NULL DEFAULT 'activity', title TEXT NOT NULL, message TEXT NOT NULL,
                is_read BOOLEAN NOT NULL DEFAULT FALSE, created_at TIMESTAMPTZ DEFAULT NOW())""")
            cur.execute("CREATE INDEX IF NOT EXISTS notifications_recipient_idx ON notifications(recipient_id,is_read,created_at DESC)")
            cur.execute("CREATE INDEX IF NOT EXISTS notifications_created_idx ON notifications(created_at DESC)")

            # Internal 1-to-1 chat tables. Existing CRM tables/features are untouched.
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_conversations (
                id SERIAL PRIMARY KEY,
                user1_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                user2_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                last_message_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                CONSTRAINT chat_conversations_users_order CHECK (user1_id < user2_id),
                CONSTRAINT chat_conversations_unique_pair UNIQUE (user1_id, user2_id)
            )""")
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_messages (
                id SERIAL PRIMARY KEY,
                conversation_id INTEGER NOT NULL REFERENCES chat_conversations(id) ON DELETE CASCADE,
                sender_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                body TEXT NOT NULL,
                is_read BOOLEAN NOT NULL DEFAULT FALSE,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                edited_at TIMESTAMPTZ,
                deleted_at TIMESTAMPTZ,
                reply_to_id INTEGER REFERENCES chat_messages(id) ON DELETE SET NULL
            )""")
            cur.execute("ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS edited_at TIMESTAMPTZ")
            cur.execute("ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS deleted_at TIMESTAMPTZ")
            cur.execute("ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS reply_to_id INTEGER REFERENCES chat_messages(id) ON DELETE SET NULL")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_conversations_user1_idx ON chat_conversations(user1_id,last_message_at DESC)")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_conversations_user2_idx ON chat_conversations(user2_id,last_message_at DESC)")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_messages_conversation_idx ON chat_messages(conversation_id,created_at ASC,id ASC)")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_messages_unread_idx ON chat_messages(conversation_id,is_read,created_at ASC)")
            # Shared Team Room: one common room for the whole approved team.
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_team_room (
                id INTEGER PRIMARY KEY DEFAULT 1,
                room_name TEXT NOT NULL DEFAULT 'Team Room',
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                CONSTRAINT chat_team_room_singleton CHECK (id=1)
            )""")
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_team_messages (
                id SERIAL PRIMARY KEY,
                room_id INTEGER NOT NULL DEFAULT 1 REFERENCES chat_team_room(id) ON DELETE CASCADE,
                sender_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                body TEXT NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                edited_at TIMESTAMPTZ,
                deleted_at TIMESTAMPTZ,
                reply_to_id INTEGER REFERENCES chat_team_messages(id) ON DELETE SET NULL
            )""")
            cur.execute("ALTER TABLE chat_team_messages ADD COLUMN IF NOT EXISTS edited_at TIMESTAMPTZ")
            cur.execute("ALTER TABLE chat_team_messages ADD COLUMN IF NOT EXISTS deleted_at TIMESTAMPTZ")
            cur.execute("ALTER TABLE chat_team_messages ADD COLUMN IF NOT EXISTS reply_to_id INTEGER REFERENCES chat_team_messages(id) ON DELETE SET NULL")
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_team_reads (
                room_id INTEGER NOT NULL DEFAULT 1 REFERENCES chat_team_room(id) ON DELETE CASCADE,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                last_read_message_id INTEGER,
                last_read_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                PRIMARY KEY (room_id,user_id)
            )""")
            cur.execute("INSERT INTO chat_team_room(id,room_name) VALUES(1,'Team Room') ON CONFLICT(id) DO NOTHING")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_team_messages_room_idx ON chat_team_messages(room_id,created_at ASC,id ASC)")
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_presence (
                user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
                last_seen TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )""")
            # Reactions: create the current schema and safely migrate older MVP databases.
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_message_reactions (
                id SERIAL PRIMARY KEY,
                private_message_id INTEGER REFERENCES chat_messages(id) ON DELETE CASCADE,
                team_message_id INTEGER REFERENCES chat_team_messages(id) ON DELETE CASCADE,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                reaction TEXT NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )""")
            # Older versions may already have chat_message_reactions without the target columns.
            # Add them before any SELECT/INSERT so existing installations do not crash.
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS id BIGSERIAL")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS private_message_id INTEGER")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS team_message_id INTEGER")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS user_id INTEGER")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS reaction TEXT")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS created_at TIMESTAMPTZ DEFAULT NOW()")
            # Preserve legacy private-message reactions when an old generic message_id column exists.
            cur.execute("""DO $$
            BEGIN
                IF EXISTS (SELECT 1 FROM information_schema.columns WHERE table_name='chat_message_reactions' AND column_name='message_id')
                   AND EXISTS (SELECT 1 FROM information_schema.columns WHERE table_name='chat_message_reactions' AND column_name='private_message_id')
                THEN
                    EXECUTE 'UPDATE chat_message_reactions SET private_message_id=message_id WHERE private_message_id IS NULL AND message_id IS NOT NULL';
                END IF;
            END $$""")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_presence_last_seen_idx ON chat_presence(last_seen)")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_reactions_private_idx ON chat_message_reactions(private_message_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_reactions_team_idx ON chat_message_reactions(team_message_id)")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS chat_reaction_private_unique_idx ON chat_message_reactions(private_message_id,user_id,reaction) WHERE private_message_id IS NOT NULL")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS chat_reaction_team_unique_idx ON chat_message_reactions(team_message_id,user_id,reaction) WHERE team_message_id IS NOT NULL")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_team_reads_user_idx ON chat_team_reads(user_id,room_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS leads_followup_idx ON leads(next_followup_date)")
            cur.execute("CREATE INDEX IF NOT EXISTS activity_logs_lead_created_idx ON activity_logs(lead_id,created_at DESC)")
            cur.execute("CREATE INDEX IF NOT EXISTS call_logs_lead_date_idx ON call_logs(lead_id,call_date DESC)")            # Ensure legacy leads remain visible to their original creator.
            cur.execute("UPDATE leads SET assigned_to=user_id WHERE assigned_to IS NULL AND user_id IS NOT NULL")
            # Never auto-promote signup users; Admin is created explicitly.
            cur.execute("UPDATE users SET role='Agent' WHERE COALESCE(role,'Agent') NOT IN ('Admin','Agent')")
            cur.execute("UPDATE users SET approval_status='Approved' WHERE approval_status IS NULL OR approval_status=''")
            cur.execute("UPDATE users SET auth_provider='password' WHERE auth_provider IS NULL OR auth_provider=''")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS users_firebase_uid_uq ON users(firebase_uid) WHERE firebase_uid IS NOT NULL")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS users_phone_uq ON users(phone) WHERE phone IS NOT NULL")
            cur.execute("ALTER TABLE users ALTER COLUMN email DROP NOT NULL")

        conn.commit()



def ensure_chat_reaction_schema():
    """Run the chat-reaction migration in its own transaction.

    This is intentionally separate from the larger init_db transaction so a
    failure in an unrelated legacy migration cannot roll back these columns.
    """
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""CREATE TABLE IF NOT EXISTS chat_message_reactions (
                id SERIAL PRIMARY KEY,
                private_message_id INTEGER,
                team_message_id INTEGER,
                user_id INTEGER,
                reaction TEXT,
                created_at TIMESTAMPTZ DEFAULT NOW()
            )""")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS private_message_id INTEGER")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS team_message_id INTEGER")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS user_id INTEGER")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS reaction TEXT")
            cur.execute("ALTER TABLE chat_message_reactions ADD COLUMN IF NOT EXISTS created_at TIMESTAMPTZ DEFAULT NOW()")
            # Older MVP builds used a generic message_id for private reactions.
            cur.execute("""DO $$
            BEGIN
                IF EXISTS (SELECT 1 FROM information_schema.columns WHERE table_name='chat_message_reactions' AND column_name='message_id')
                   AND EXISTS (SELECT 1 FROM information_schema.columns WHERE table_name='chat_message_reactions' AND column_name='private_message_id')
                THEN
                    EXECUTE 'UPDATE chat_message_reactions SET private_message_id=message_id WHERE private_message_id IS NULL AND message_id IS NOT NULL';
                END IF;
            END $$""")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_reactions_private_idx ON chat_message_reactions(private_message_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS chat_reactions_team_idx ON chat_message_reactions(team_message_id)")
        conn.commit()


@app.on_event("startup")

def startup():

    try:

        init_db()

        print("Database initialization complete.")

    except Exception as exc:

        print(f"Database initialization error: {exc}")

    # Run this migration separately so unrelated legacy DB errors cannot roll it back.
    try:
        ensure_chat_reaction_schema()
        print("Chat reaction schema migration complete.")
    except Exception as exc:
        print(f"Chat reaction schema migration error: {exc}")

# ============================================================

# Passwords / JWT

# ============================================================

PBKDF2_ITERATIONS = 260_000

def hash_password(password: str, salt: Optional[bytes] = None):

    salt = salt or secrets.token_bytes(16)

    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)

    return base64.b64encode(digest).decode("ascii"), base64.b64encode(salt).decode("ascii")



def verify_password(password: str, password_hash: str, salt_b64: str) -> bool:

    try:

        salt = base64.b64decode(salt_b64)

        expected = base64.b64decode(password_hash)

        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)

        return hmac.compare_digest(actual, expected)

    except Exception:

        return False



def create_token(user_id: int, email: str):

    now = datetime.now(timezone.utc)

    return jwt.encode({"sub": str(user_id), "email": email, "iat": now, "exp": now + timedelta(hours=JWT_EXPIRE_HOURS)}, JWT_SECRET, algorithm=JWT_ALGORITHM)



def decode_token(token: str):

    try:

        return jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])

    except jwt.ExpiredSignatureError:

        raise HTTPException(status_code=401, detail="Session expired. Please log in again.")

    except jwt.InvalidTokenError:

        raise HTTPException(status_code=401, detail="Invalid authentication token.")



def get_current_user(authorization: Optional[str] = Header(None)):

    if not authorization or not authorization.startswith("Bearer "):

        raise HTTPException(status_code=401, detail="Authentication required.")

    payload = decode_token(authorization.split(" ", 1)[1])

    try:

        user_id = int(payload["sub"])

    except (KeyError, ValueError, TypeError):

        raise HTTPException(status_code=401, detail="Invalid authentication token.")

    with db_conn() as conn:

        with conn.cursor() as cur:

            cur.execute("SELECT id, email, name, COALESCE(role, 'Agent'), team_leader_id, COALESCE(active, TRUE), COALESCE(approval_status, 'Approved'), phone, COALESCE(auth_provider, 'password') FROM users WHERE id=%s", (user_id,))

            user = cur.fetchone()

    if not user:

        raise HTTPException(status_code=401, detail="User not found.")

    if not user[5]:
        raise HTTPException(status_code=403, detail="This account is disabled. Contact the Admin.")
    if user[6] != "Approved" and user[3] != "Admin":
        raise HTTPException(status_code=403, detail="Your account is pending Admin approval.")
    return {"id": user[0], "email": user[1], "name": user[2], "role": user[3], "team_leader_id": user[4], "approval_status": user[6], "phone": user[7], "auth_provider": user[8]}

# ============================================================

# Models

# ============================================================

class SignupRequest(BaseModel):

    email: str

    password: str

    name: Optional[str] = None

class LoginRequest(BaseModel):

    email: str

    password: str

class FirebaseAuthRequest(BaseModel):
    id_token: str = Field(min_length=20, max_length=10000)
    name: Optional[str] = Field(None, max_length=200)
    phone: Optional[str] = Field(None, max_length=50)
    auth_provider: str = Field(default="firebase", max_length=30)

class AdminAgentCreateRequest(BaseModel):
    email: str
    password: str = Field(min_length=6, max_length=200)
    name: Optional[str] = Field(None, max_length=200)

class AdminAgentUpdateRequest(BaseModel):
    email: Optional[str] = None
    name: Optional[str] = Field(None, max_length=200)
    password: Optional[str] = Field(None, min_length=6, max_length=200)
    active: Optional[bool] = None

class AdminResetPasswordRequest(BaseModel):
    new_password: str = Field(min_length=6, max_length=200)

class SearchRequest(BaseModel):

    city: str = Field(min_length=2, max_length=120)

    category: Optional[str] = None

    keyword: Optional[str] = None

    limit: int = Field(default=10, ge=1, le=100)

    city_wide: bool = False

    radius_meters: int = Field(default=SEARCH_RADIUS_METERS, ge=250, le=5000)

    enrich: bool = False

class StatusUpdateRequest(BaseModel):

    status: str

class NotesUpdateRequest(BaseModel):

    notes: Optional[str] = None

class LeadUpdateRequest(BaseModel):
    name: Optional[str] = Field(None, min_length=1, max_length=250)
    address: Optional[str] = Field(None, max_length=1000)
    city: Optional[str] = Field(None, max_length=120)
    region: Optional[str] = Field(None, max_length=120)
    country: Optional[str] = Field(None, max_length=120)
    phone: Optional[str] = Field(None, max_length=100)
    website: Optional[str] = Field(None, max_length=500)
    email: Optional[str] = Field(None, max_length=320)
    instagram: Optional[str] = Field(None, max_length=500)
    facebook: Optional[str] = Field(None, max_length=500)
    twitter: Optional[str] = Field(None, max_length=500)
    category: Optional[str] = Field(None, max_length=150)
    notes: Optional[str] = Field(None, max_length=5000)
    status: Optional[str] = None
    assigned_to: Optional[int] = None
    last_contacted_date: Optional[datetime] = None
    next_followup_date: Optional[datetime] = None

class AssignmentRequest(BaseModel):
    assigned_to: Optional[int] = None

class FollowupRequest(BaseModel):
    next_followup_date: Optional[datetime] = None

class CallLogRequest(BaseModel):
    call_date: Optional[datetime] = None
    call_outcome: str = Field(min_length=1, max_length=100)
    notes: Optional[str] = Field(None, max_length=5000)

class BulkLeadActionRequest(BaseModel):
    lead_ids: List[int] = Field(min_length=1, max_length=500)
    action: str
    status: Optional[str] = None
    assigned_to: Optional[int] = None

class GroupCreateRequest(BaseModel):

    name: str

class GroupBulkAddRequest(BaseModel):

    lead_ids: List[int]

class ChatStartRequest(BaseModel):
    user_id: int = Field(gt=0)

class ChatMessageRequest(BaseModel):
    body: str = Field(min_length=1, max_length=5000)
    reply_to_id: Optional[int] = Field(None, gt=0)

class ChatEditMessageRequest(BaseModel):
    body: str = Field(min_length=1, max_length=5000)

class ChatReactionRequest(BaseModel):
    reaction: str = Field(min_length=1, max_length=16)

# ============================================================

# Auth endpoints

# ============================================================

@app.post("/api/signup")

def signup(req: SignupRequest):

    email = req.email.strip().lower()

    name = (req.name or "").strip()

    if "@" not in email or "." not in email.split("@")[-1]:

        raise HTTPException(status_code=400, detail="Please enter a valid email address.")

    if len(req.password) < 6:

        raise HTTPException(status_code=400, detail="Password must be at least 6 characters.")

    password_hash, salt = hash_password(req.password)

    try:

        with db_conn() as conn:

            with conn.cursor() as cur:

                cur.execute("SELECT id FROM users WHERE email=%s", (email,))

                if cur.fetchone():

                    raise HTTPException(status_code=400, detail="An account with this email already exists.")

                role = "Agent"
                cur.execute("INSERT INTO users(email,name,password_hash,salt,role,active,approval_status,auth_provider) VALUES(%s,%s,%s,%s,%s,FALSE,'Pending','password') RETURNING id", (email, name or None, password_hash, salt, role))

                user_id = cur.fetchone()[0]

            conn.commit()

    except HTTPException:

        raise

    except Exception as exc:

        print(f"Signup database error: {exc}")

        raise HTTPException(status_code=500, detail=f"Signup database error: {type(exc).__name__}.")

    return {"pending": True, "message": "Account created. Wait for Admin approval before logging in.", "user": {"id": user_id, "email": email, "name": name, "role": role, "approval_status": "Pending"}}

@app.post("/api/login")

def login(req: LoginRequest):

    email = req.email.strip().lower()

    try:

        with db_conn() as conn:

            with conn.cursor() as cur:

                cur.execute("SELECT id,email,name,password_hash,salt,COALESCE(role, 'Agent'),team_leader_id,COALESCE(active, TRUE),COALESCE(approval_status, 'Approved'),phone,COALESCE(auth_provider, 'password') FROM users WHERE email=%s", (email,))

                user = cur.fetchone()

    except Exception as exc:

        raise HTTPException(status_code=500, detail=f"Login database error: {type(exc).__name__}.")

    if not user or not verify_password(req.password, user[3], user[4]):

        raise HTTPException(status_code=401, detail="Invalid email or password.")
    if not user[7]:
        raise HTTPException(status_code=403, detail="This account is disabled. Contact the Admin.")
    if user[8] != "Approved" and user[5] != "Admin":
        raise HTTPException(status_code=403, detail="Your account is pending Admin approval.")

    with db_conn() as conn:
        with conn.cursor() as cur:
            log_activity(cur, user[0], None, "logged_in", {})
        conn.commit()
    return {"token": create_token(user[0], user[1]), "user": {"id": user[0], "email": user[1], "name": user[2], "role": user[5], "team_leader_id": user[6], "approval_status": user[8], "phone": user[9], "auth_provider": user[10]}}

def verify_firebase_id_token(id_token: str) -> dict:
    if not FIREBASE_API_KEY:
        raise HTTPException(status_code=503, detail="Firebase authentication is not configured on the server.")
    try:
        r = requests.post(
            "https://identitytoolkit.googleapis.com/v1/accounts:lookup",
            params={"key": FIREBASE_API_KEY},
            json={"idToken": id_token},
            timeout=15,
        )
    except requests.RequestException:
        raise HTTPException(status_code=502, detail="Firebase authentication service is temporarily unavailable.")
    if r.status_code != 200:
        raise HTTPException(status_code=401, detail="Firebase authentication failed. Please sign in again.")
    try:
        users = r.json().get("users") or []
    except Exception:
        users = []
    if not users:
        raise HTTPException(status_code=401, detail="Firebase account could not be verified.")
    return users[0]

@app.post("/api/auth/firebase-sync")
def firebase_sync(req: FirebaseAuthRequest):
    info = verify_firebase_id_token(req.id_token)
    firebase_uid = str(info.get("localId") or "").strip()
    email = (info.get("email") or "").strip().lower() or None
    phone = (info.get("phoneNumber") or req.phone or "").strip() or None
    name = (req.name or info.get("displayName") or "").strip() or None
    provider = (req.auth_provider or "firebase").strip().lower()
    if not firebase_uid:
        raise HTTPException(status_code=401, detail="Firebase account is missing a user ID.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""SELECT id,email,name,COALESCE(role,'Agent'),COALESCE(active,TRUE),COALESCE(approval_status,'Approved'),phone
                           FROM users WHERE firebase_uid=%s OR lower(email)=lower(%s::text) OR phone=%s
                           ORDER BY CASE WHEN COALESCE(role,'Agent')='Admin' THEN 0 ELSE 1 END LIMIT 1""", (firebase_uid,email,phone))
            existing = cur.fetchone()
            if existing and existing[3] == "Admin":
                raise HTTPException(status_code=403, detail="The Admin account uses the secure email/password login.")
            if existing:
                uid = existing[0]
                if not existing[4]:
                    raise HTTPException(status_code=403, detail="This account is disabled. Contact the Admin.")
                approval = existing[5]
                cur.execute("UPDATE users SET firebase_uid=%s,phone=COALESCE(%s,phone),auth_provider=%s,name=COALESCE(%s,name) WHERE id=%s", (firebase_uid,phone,provider,name,uid))
            else:
                cur.execute("""INSERT INTO users(email,name,password_hash,salt,role,active,firebase_uid,phone,approval_status,auth_provider)
                               VALUES(%s,%s,'','', 'Agent', FALSE,%s,%s,'Pending',%s) RETURNING id""", (email,name,firebase_uid,phone,provider))
                uid = cur.fetchone()[0]
                approval = "Pending"
                log_activity(cur,uid,None,"account_pending_approval",{"auth_provider":provider})
            conn.commit()
    if approval != "Approved":
        raise HTTPException(status_code=403, detail="Account created successfully. Wait for Admin approval before logging in.")
    login_email = email or f"firebase:{firebase_uid}"
    with db_conn() as conn:
        with conn.cursor() as cur:
            log_activity(cur, uid, None, "logged_in", {"auth_provider": provider})
        conn.commit()
    return {"token": create_token(uid, login_email), "user": {"id":uid,"email":email,"name":name,"role":"Agent","approval_status":"Approved","phone":phone,"auth_provider":provider}}

@app.get("/api/me")

def me(current_user: dict = Depends(get_current_user)):

    return {"user": current_user}

@app.get("/api/categories")

def categories():

    return {"categories": list(CATEGORY_MAP.keys())}

# ============================================================

# SerpApi Google Maps provider (primary local-business source)

# ============================================================

def serpapi_get(params: dict, timeout: int = 30):

    if not SERPAPI_API_KEY:

        raise ProviderError("SERPAPI_API_KEY is not configured.")

    query = dict(params)

    query.update({"api_key": SERPAPI_API_KEY})
    query.setdefault("engine", "google_maps")

    try:

        response = requests.get("https://serpapi.com/search.json", params=query, timeout=timeout)

    except requests.RequestException:

        raise ProviderError("SerpApi is temporarily unavailable. Please try again later.")

    if response.status_code == 401:

        raise ProviderError("SerpApi API key is invalid or unauthorized.")

    if response.status_code == 429:

        raise ProviderError("SerpApi rate limit reached. Please try again later.")

    if response.status_code != 200:

        try:

            detail = response.json().get("error")

        except Exception:

            detail = None

        raise ProviderError(f"SerpApi returned HTTP {response.status_code}" + (f": {detail}" if detail else "."))

    try:

        data = response.json()

    except ValueError:

        raise ProviderError("SerpApi returned an invalid response.")

    if data.get("error"):

        raise ProviderError(str(data["error"]))

    return data



def search_serpapi_maps(lat: Optional[float], lon: Optional[float], query: str, limit: int):

    """Search Google Maps through SerpApi. Coordinates are optional; when

    available they improve geographic targeting and enable pagination."""

    target = min(max(int(limit), 1), MAX_LIMIT)

    results = []

    # SerpApi documents ll + start for paginated Maps searches. Without a

    # reliable geocoded center, use the first Maps page and rely on the city

    # included in q instead of requiring the old LatLng provider.

    max_pages = min(5, (target + 19) // 20) if lat is not None and lon is not None else 1

    for page in range(max_pages):

        params = {

            "q": query,

            "type": "search",

            "hl": "en",

            "gl": "in",

        }

        if lat is not None and lon is not None:

            params["ll"] = f"@{float(lat)},{float(lon)},14z"

            params["start"] = page * 20

        data = serpapi_get(params)

        page_results = data.get("local_results") or []

        if not isinstance(page_results, list):

            raise ProviderError("SerpApi returned an unexpected Google Maps response.")

        results.extend(page_results)

        if len(results) >= target or len(page_results) < 20:

            break

    return results[:target]



def extract_serpapi_lead(

    place: dict,

    requested_category: Optional[str],

    resolved_city: str,

    center_lat: float,

    center_lon: float,

    radius_meters: int = SEARCH_RADIUS_METERS,

):

    name = place.get("title")

    gps = place.get("gps_coordinates") or {}

    lat, lon = gps.get("latitude"), gps.get("longitude")

    if not name or lat is None or lon is None:

        return None

    try:

        lat, lon = float(lat), float(lon)

    except (ValueError, TypeError):

        return None

    # Google Maps results are not guaranteed to stay inside ll, so enforce

    # our own geographic boundary.

    if center_lat is not None and center_lon is not None:

        if distance_km(center_lat, center_lon, lat, lon) > radius_meters / 1000:

            return None

    provider_id = place.get("place_id") or place.get("data_cid") or place.get("data_id")

    place_id = str(provider_id or f"{name}|{lat:.6f}|{lon:.6f}")

    links = place.get("links") or {}

    website = place.get("website") or links.get("website")

    maps_url = (

        links.get("directions")

        or "https://www.google.com/maps/search/?api=1&query="

        + quote_plus(f"{name}, {resolved_city}")

    )

    return {

        "place_id": place_id,

        "provider_place_id": str(provider_id) if provider_id else None,

        "name": str(name),

        "address": place.get("address"),

        "city": resolved_city,

        "region": None,

        "country": place.get("country") or "India",

        "latitude": lat,

        "longitude": lon,

        "phone": place.get("phone"),

        "website": website,

        "email": place.get("email"),

        "instagram": place.get("instagram"),

        "facebook": place.get("facebook"),

        "twitter": place.get("twitter"),

        "category": requested_category or place.get("type"),

        "rating": place.get("rating"),

        "reviews": place.get("reviews"),

        "maps_url": maps_url,

        "source": "Google Maps via SerpApi",

        "sources": ["Google Maps via SerpApi"],

        "provider_category": place.get("type"),

        "distance_m": place.get("distance"),

        "match_score": None,

    }



# ============================================================

# LatLng provider

# ============================================================

class ProviderError(Exception):

    pass



def latlng_headers():

    if not LATLNG_API_KEY:

        raise ProviderError("LATLNG_API_KEY is not configured.")

    return {"X-Api-Key": LATLNG_API_KEY, "Accept": "application/json"}



def _latlng_get(url: str, params: dict, timeout: int = 20):

    try:

        response = requests.get(url, params=params, headers=latlng_headers(), timeout=timeout)

    except requests.RequestException:

        raise ProviderError("LatLng API is temporarily unavailable. Please try again later.")

    if response.status_code == 429:

        raise ProviderError("LatLng rate limit reached. Please try again later.")

    if response.status_code == 401:

        raise ProviderError("LatLng API key is invalid or not authorized.")

    if response.status_code != 200:

        raise ProviderError(f"LatLng API returned HTTP {response.status_code}.")

    try:

        return response.json()

    except ValueError:

        raise ProviderError("LatLng API returned an invalid response.")



def resolve_city(city: str):

    """Resolve the requested city without making LatLng mandatory.

    SerpApi is the primary provider. If its city lookup does not expose a

    reliable coordinate, the search can still proceed using the city name in

    the Google Maps query. LatLng is only an optional coordinate fallback.

    """

    raw = city.strip()

    city_name = raw.split(",", 1)[0].strip()

    if not city_name:

        raise ProviderError("City is required.")

    if SERPAPI_API_KEY:

        try:

            data = serpapi_get({

                "q": raw,

                "type": "search",

                "hl": "en",

                "gl": "in",

            })

            results = data.get("local_results") or []

            if isinstance(results, list):

                target = city_name.casefold()

                best = None

                for item in results:

                    title = str(item.get("title") or "").strip()

                    address = str(item.get("address") or "").casefold()

                    if title.casefold() == target or target in address:

                        best = item

                        break

                # Do not treat an arbitrary business as the city center.

                # If Maps does not return a city-level result, search by the

                # city name without a coordinate filter.

                gps = (best or {}).get("gps_coordinates") or {}

                lat = gps.get("latitude")

                lon = gps.get("longitude")

                try:

                    lat = float(lat) if lat is not None else None

                    lon = float(lon) if lon is not None else None

                except (ValueError, TypeError):

                    lat, lon = None, None

                return {

                    "lat": lat,

                    "lon": lon,

                    "resolved_city": city_name,

                    "state": "",

                    "country": "India",

                }

        except ProviderError:

            # If LatLng is configured, it remains a safe coordinate fallback.

            if not LATLNG_API_KEY:

                raise

    if LATLNG_API_KEY:

        data = _latlng_get(LATLNG_GEOCODE_URL, {"q": raw, "limit": 10, "lang": "en"}, 15)

        features = data.get("features", [])

        target = city_name.casefold()

        for feature in features:

            props = feature.get("properties") or {}

            coords = (feature.get("geometry") or {}).get("coordinates") or []

            if len(coords) < 2:

                continue

            candidates = {str(props.get(k) or "").strip().casefold() for k in ("city", "town", "municipality", "village", "name")}

            candidates.discard("")

            country = str(props.get("countrycode") or props.get("country_code") or "").strip().casefold()

            if target not in candidates and city_name.casefold() not in str(props.get("display_name") or "").casefold():

                continue

            if country and country != "in":

                continue

            try:

                lon, lat = float(coords[0]), float(coords[1])

            except (ValueError, TypeError):

                continue

            return {

                "lat": lat,

                "lon": lon,

                "resolved_city": str(props.get("city") or props.get("town") or props.get("municipality") or props.get("village") or props.get("name") or city_name),

                "state": props.get("state") or "",

                "country": props.get("country") or "India",

            }

    raise ProviderError(f"Could not locate '{city_name}'. Check your SerpApi API key and try the city name only.")



def distance_km(lat1, lon1, lat2, lon2):

    from math import radians, sin, cos, sqrt, atan2

    r = 6371.0088

    dlat, dlon = radians(lat2 - lat1), radians(lon2 - lon1)

    a = sin(dlat / 2) ** 2 + cos(radians(lat1)) * cos(radians(lat2)) * sin(dlon / 2) ** 2

    return 2 * r * atan2(sqrt(a), sqrt(1 - a))



def search_nearby(lat, lon, category, limit, radius=SEARCH_RADIUS_METERS):

    """Search businesses around a coordinate using LatLng Places Nearby."""

    latlng_category = CATEGORY_MAP.get(category)

    if not latlng_category:

        raise ProviderError(f"Category '{category}' is not supported.")

    params = {

        "lat": float(lat),

        "lon": float(lon),

        "radius": min(max(int(radius), 250), 5000),

        "category": latlng_category,

        "country": "IN",

        "limit": min(max(int(limit), 1), 50),

    }

    data = _latlng_get(LATLNG_PLACES_NEARBY_URL, params, 20)

    places = data.get("places", []) if isinstance(data, dict) else []

    if not isinstance(places, list):

        raise ProviderError("LatLng returned an unexpected Places response.")

    return places

def search_keyword(lat, lon, keyword, limit):

    q = keyword.strip()

    if len(q) < 2:

        raise ProviderError("Keyword must contain at least 2 characters.")

    data = _latlng_get(LATLNG_PLACES_SEARCH_URL, {"q": q, "lat": lat, "lon": lon, "limit": min(max(limit, 1), 50)}, 20)

    return data.get("places", [])



def city_grid(location: dict, enabled: bool, spacing_meters: Optional[int] = None):
    if not enabled or location.get("lat") is None or location.get("lon") is None:
        return [location]
    from math import cos, radians
    spacing = spacing_meters or CITY_WIDE_GRID_SPACING_METERS
    offsets = [(0,0),(0,1),(1,0),(0,-1),(-1,0),(1,1),(1,-1),(-1,1),(-1,-1)][:CITY_WIDE_MAX_ZONES]
    lat_delta = spacing / 111_320
    lon_delta = spacing / max(1, 111_320 * abs(cos(radians(location["lat"]))))
    return [{**location, "lat": location["lat"] + y * lat_delta, "lon": location["lon"] + x * lon_delta} for x, y in offsets]



def extract_lead(place: dict, requested_category: Optional[str], resolved_city: str, center_lat: float, center_lon: float, radius_meters: int = SEARCH_RADIUS_METERS):

    name, lat, lon = place.get("name"), place.get("lat"), place.get("lon")

    if not name or lat is None or lon is None:

        return None

    try:

        lat, lon = float(lat), float(lon)

    except (ValueError, TypeError):

        return None

    # Keep results geographically relevant to the requested city/grid zone.

    if distance_km(center_lat, center_lon, lat, lon) > radius_meters / 1000:

        return None

    provider_id = place.get("place_id") or place.get("id")

    place_id = str(provider_id or f"{name}|{lat:.6f}|{lon:.6f}")

    maps_url = "https://www.google.com/maps/search/?api=1&query=" + quote_plus(f"{name}, {resolved_city}")

    return {

        "place_id": place_id,

        "provider_place_id": str(provider_id) if provider_id else None,

        "name": str(name),

        "address": place.get("address") or place.get("formatted_address"),

        "city": resolved_city,

        "region": None,

        "country": "India",

        "latitude": lat,

        "longitude": lon,

        "phone": place.get("phone"),

        "website": place.get("website"),

        "email": place.get("email"),

        "instagram": place.get("instagram"),

        "facebook": place.get("facebook"),

        "twitter": place.get("twitter"),

        "category": requested_category or place.get("category"),

        "rating": place.get("rating"),

        "reviews": place.get("reviews"),

        "maps_url": maps_url,

        "source": "LatLng Places",

        "sources": ["LatLng Places"],

        "provider_category": place.get("category"),

        "distance_m": place.get("distance_m"),

        "match_score": place.get("match_score"),

    }

# ============================================================

# Dedupe / merge

# ============================================================

def norm(value):

    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())



def phone_key(value):

    return re.sub(r"\D", "", str(value or ""))[-10:] if value else ""



def domain_key(value):

    if not value:

        return ""

    try:

        host = urlparse(value if "//" in value else "https://" + value).netloc.lower().split(":")[0]

        return host[4:] if host.startswith("www.") else host

    except Exception:

        return ""



def dedupe_leads(leads: list[dict], limit: int):

    groups = []

    indexes = {}

    for lead in leads:

        keys = [

            ("id", norm(lead.get("provider_place_id") or lead.get("place_id"))),

            ("phone", phone_key(lead.get("phone"))),

            ("domain", domain_key(lead.get("website"))),

            ("nameaddr", norm((lead.get("name") or "") + " " + (lead.get("address") or ""))),

        ]

        found = None

        for kind, key in keys:

            if key and (kind, key) in indexes:

                found = indexes[(kind, key)]

                break

        if found is None:

            found = len(groups)

            groups.append(dict(lead))

        else:

            target = groups[found]

            for key, value in lead.items():

                if key in {"sources", "source"}:

                    continue

                if not target.get(key) and value not in (None, ""):

                    target[key] = value

            target["sources"] = list(dict.fromkeys([*(target.get("sources") or []), *(lead.get("sources") or []), lead.get("source") or ""]))

            target["sources"] = [x for x in target["sources"] if x]

            target["source"] = ", ".join(target["sources"])

        for kind, key in keys:

            if key:

                indexes[(kind, key)] = found

        if len(groups) >= limit and len(groups) >= limit * 2:

            break

    return groups[:limit]

# ============================================================

# Public website enrichment (conservative)

# ============================================================

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)

SOCIAL_RE = {

    "instagram": re.compile(r"https?://(?:www\.)?instagram\.com/[^\"'\s<>]+", re.I),

    "facebook": re.compile(r"https?://(?:www\.)?facebook\.com/[^\"'\s<>]+", re.I),

    "twitter": re.compile(r"https?://(?:www\.)?(?:twitter\.com|x\.com)/[^\"'\s<>]+", re.I),
    "linkedin": re.compile(r"https?://(?:www\.)?linkedin\.com/(?:company|showcase)/[^\"'\s<>]+", re.I),
    "youtube": re.compile(r"https?://(?:www\.)?(?:youtube\.com|youtu\.be)/[^\"'\s<>]+", re.I),
    "tiktok": re.compile(r"https?://(?:www\.)?tiktok\.com/@?[^\"'\s<>]+", re.I),

}



def enrich_public_website(url: Optional[str]):

    if not url or not str(url).startswith(("http://", "https://")):

        return {}

    try:

        rp = robotparser.RobotFileParser()

        base = f"{urlparse(url).scheme}://{urlparse(url).netloc}"

        rp.set_url(base.rstrip("/") + "/robots.txt")

        try:

            rp.read()

            if not rp.can_fetch("LeadSearchingEngine/1.0", url):

                return {}

        except Exception:

            # A failed robots fetch is not treated as permission to crawl; skip enrichment.

            return {}

        response = requests.get(url, headers={"User-Agent": "LeadSearchingEngine/1.0"}, timeout=ENRICH_TIMEOUT, allow_redirects=True)

        if response.status_code != 200 or "text/html" not in response.headers.get("content-type", "").lower():

            return {}

        text = response.text[:1_000_000]

        result = {"sources": ["Business Website Enrichment"], "enrichment_source": response.url, "enrichment_at": datetime.now(timezone.utc).isoformat()}

        email = EMAIL_RE.search(text)

        if email:

            result["email"] = email.group(0)

        for key, pattern in SOCIAL_RE.items():

            match = pattern.search(text)

            if match:

                result[key] = match.group(0).rstrip(".,);'\"")

        return result

    except Exception:

        return {}

# ============================================================

# AI analysis

# ============================================================

def fallback_analysis(lead: dict):

    phone, website = bool(lead.get("phone")), bool(lead.get("website"))

    social = bool(lead.get("instagram") or lead.get("facebook") or lead.get("twitter"))

    score = 45 + (10 if phone else 0) + (10 if website else 0) + (5 if social else 0)

    if lead.get("rating") is not None: score += 5

    if lead.get("reviews") is not None and lead.get("reviews") >= 20: score += 5

    score = min(score, 80)

    if not phone and not website and not social:

        lead_type, recommendation = "Needs More Information", "Verify contact details manually before outreach."

    elif phone and not website:

        lead_type, recommendation = "Medium Potential", "Consider contacting the business about its website or online presence."

    else:

        lead_type, recommendation = "Needs More Information", "Review the available business signals before prioritizing outreach."

    return {"lead_score": score, "lead_type": lead_type, "ai_recommendation": recommendation,

            "ai_reason": "Conservative assessment based only on fields returned by the authorized data source. Missing fields do not prove the business lacks them."}



def score_to_lead_priority(score: int) -> str:
    score = max(0, min(100, int(score)))
    if score >= 70:
        return "Hot Lead"
    if score >= 45:
        return "Warm Lead"
    return "Cold Lead"

def gemini_analysis(lead: dict):
    fallback = fallback_analysis(lead)
    if not GEMINI_API_KEY:
        fallback["lead_type"] = score_to_lead_priority(fallback["lead_score"])
        return fallback
    prompt = f'''Analyze this business lead for a digital marketing/web-services agency.\nBusiness: {lead.get("name")}\nCategory: {lead.get("category")}\nPhone: {"yes" if lead.get("phone") else "no"}\nWebsite: {"yes" if lead.get("website") else "no"}\nSocial: {"yes" if (lead.get("instagram") or lead.get("facebook") or lead.get("twitter")) else "no"}\nRating: {lead.get("rating", "not available")}\nReviews: {lead.get("reviews", "not available")}\nIMPORTANT: Missing fields only mean the source did not return that field. Do not invent facts. Return ONLY JSON with lead_score and reason.'''
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={GEMINI_API_KEY}"
    try:
        response = requests.post(url, json={"contents":[{"parts":[{"text":prompt}]}]}, timeout=12)
        if response.status_code != 200:
            fallback["lead_type"] = score_to_lead_priority(fallback["lead_score"]); return fallback
        text = response.json()["candidates"][0]["content"]["parts"][0]["text"].strip()
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I)
        parsed=json.loads(text)
        score=max(0,min(100,int(parsed.get("lead_score",fallback["lead_score"]))))
        return {"lead_score":score,"lead_type":score_to_lead_priority(score),"ai_recommendation":fallback["ai_recommendation"],"ai_reason":str(parsed.get("reason") or fallback["ai_reason"])[:1000]}
    except Exception:
        fallback["lead_type"] = score_to_lead_priority(fallback["lead_score"]); return fallback
# Search

# ============================================================

def _check_search_rate(user_id: int):
    now = datetime.now(timezone.utc).timestamp()
    events = _SEARCH_EVENTS[user_id]
    while events and now - events[0] > SEARCH_WINDOW_SECONDS:
        events.popleft()
    if len(events) >= SEARCH_MAX_REQUESTS:
        raise HTTPException(status_code=429, detail="Search limit reached temporarily. Please wait a few minutes and try again.")
    events.append(now)

@app.post("/api/search")
def search(req: SearchRequest, current_user: dict = Depends(get_current_user)):
    _check_search_rate(current_user["id"])
    city = req.city.strip()
    category = req.category.strip() if req.category else None
    keyword = req.keyword.strip() if req.keyword else None
    if not city:
        raise HTTPException(status_code=400, detail="City is required.")
    if bool(category) == bool(keyword):
        raise HTTPException(status_code=400, detail="Provide either a category or a custom keyword, not both.")
    if category and category not in CATEGORY_MAP:
        raise HTTPException(status_code=400, detail=f"Category '{category}' is not supported.")
    if keyword and (len(keyword) < 2 or len(keyword) > 100 or any(ord(c) < 32 for c in keyword)):
        raise HTTPException(status_code=400, detail="Keyword must be 2-100 characters and contain no control characters.")

    try:
        location = resolve_city(city)
    except ProviderError as exc:
        raise HTTPException(status_code=502, detail=str(exc))

    limit = min(max(int(req.limit), 1), MAX_LIMIT)
    radius_meters = min(max(int(req.radius_meters), 250), 5000)
    warnings = []
    used_sources = []
    leads = []
    search_query = f"{category} in {location['resolved_city']}" if category else f"{keyword} in {location['resolved_city']}"
    zones = city_grid(location, req.city_wide)
    zone_limit = min(20, max(1, limit))

    # Provider registry: each provider can fail independently. Results are merged
    # and deduplicated, so adding a future authorized provider only needs a new
    # provider adapter/runner rather than a rewrite of the search endpoint.
    provider_runners = []
    if SERPAPI_API_KEY:
        provider_runners.append(("Google Maps via SerpApi", "serpapi"))
    if LATLNG_API_KEY:
        provider_runners.append(("LatLng Places", "latlng"))
    if not provider_runners:
        raise HTTPException(status_code=503, detail="No lead search provider is configured.")

    for zone in zones:
        remaining = max(limit * 2 - len(leads), 1)
        per_provider = min(zone_limit, remaining)
        for source_name, provider_name in provider_runners:
            try:
                if provider_name == "serpapi":
                    places = search_serpapi_maps(zone.get("lat"), zone.get("lon"), search_query, per_provider)
                    for place in places:
                        lead = extract_serpapi_lead(
                            place, category, location["resolved_city"], zone.get("lat"), zone.get("lon"), radius_meters
                        )
                        if lead:
                            leads.append(lead)
                else:
                    if zone.get("lat") is None or zone.get("lon") is None:
                        continue
                    places = search_nearby(zone["lat"], zone["lon"], category, per_provider, radius=radius_meters) if category else search_keyword(zone["lat"], zone["lon"], keyword, per_provider)
                    for place in places:
                        lead = extract_lead(
                            place, category, location["resolved_city"], zone["lat"], zone["lon"], radius_meters
                        )
                        if lead:
                            leads.append(lead)
                if places:
                    used_sources.append(source_name)
            except ProviderError as exc:
                warnings.append(f"{source_name}: {exc}")

        # Once enough raw results exist, stop adding expensive provider calls.
        if len(leads) >= limit * 2:
            break

    leads = dedupe_leads(leads, limit)

    if not leads and warnings:
        # If one provider failed but another returned no matches, expose the
        # provider issue without leaking API keys or backend internals.
        raise HTTPException(status_code=502, detail=warnings[0])
    if not leads:
        warnings.append(f"No matching businesses were returned for {category or keyword!r} in {location['resolved_city']}.")

    if req.enrich:
        for lead in leads:
            extra = enrich_public_website(lead.get("website"))
            for key, value in extra.items():
                if key != "sources" and not lead.get(key):
                    lead[key] = value
            lead["sources"] = list(dict.fromkeys([*(lead.get("sources") or []), *(extra.get("sources") or [])]))
            lead["source"] = ", ".join(lead["sources"])
            if extra.get("sources"):
                used_sources.extend(extra["sources"])

    used_sources = list(dict.fromkeys([x for x in used_sources if x]))
    saved = []
    duplicates = []
    for lead in leads:
        lead.setdefault("enrichment_source", None)
        lead.setdefault("enrichment_at", None)

    if DATABASE_URL:
        with db_conn() as conn:
            with conn.cursor() as cur:
                for lead in leads:
                    for social_key in ("instagram", "facebook", "twitter", "linkedin", "youtube", "tiktok"):
                        lead.setdefault(social_key, None)
                    lead.setdefault("social_evidence", [])
                    lead.setdefault("social_status", "not_checked")
                    if req.enrich:
                        social = discover_public_social_profiles(lead)
                        for social_key in ("linkedin", "instagram", "facebook", "twitter", "youtube", "tiktok"):
                            if social.get(social_key):
                                lead[social_key] = social[social_key]
                        lead["social_status"] = social.get("social_status")
                        lead["social_evidence"] = social.get("social_evidence") or []
                    else:
                        lead.setdefault("social_status", "not_checked")
                        lead.setdefault("social_evidence", [])
                    lead.update(gemini_analysis(lead))
                    # Smart DB merge: provider IDs are strongest; phone/domain/name+address
                    # are used as cross-provider fallback signals. Existing status/notes survive.
                    provider_id = lead.get("provider_place_id")
                    place_id = lead.get("place_id")
                    phone = phone_key(lead.get("phone"))
                    domain = domain_key(lead.get("website"))
                    name_key = norm(lead.get("name"))
                    address_key = norm(lead.get("address"))
                    cur.execute("""SELECT id,user_id,place_id,provider_place_id,name,address,phone,website,sources,status,notes,created_at,updated_at,assigned_to
                                   FROM leads WHERE (provider_place_id=%s OR place_id=%s OR phone=%s OR lower(name)=lower(%s))
                                   ORDER BY id LIMIT 50""",
                                (provider_id, place_id, phone or None, lead.get("name")))
                    candidates = cur.fetchall()
                    existing = None
                    for cand in candidates:
                        cand_domain = domain_key(cand[7])
                        cand_name = norm(cand[4])
                        cand_addr = norm(cand[5])
                        if (provider_id and cand[3] == provider_id) or (place_id and cand[2] == place_id) or (phone and phone_key(cand[6]) == phone) or (domain and cand_domain == domain) or (name_key and cand_name == name_key and address_key and cand_addr == address_key):
                            existing = cand
                            break

                    if existing:
                        existing_id = existing[0]
                        existing_sources = existing[8] or []
                        merged_sources = list(dict.fromkeys([*(existing_sources if isinstance(existing_sources, list) else []), *(lead.get("sources") or []), lead.get("source") or ""]))
                        # Update only with available non-conflicting data; preserve CRM fields.
                        updates = {
                            "provider_place_id": lead.get("provider_place_id") or existing[3],
                            "name": lead.get("name") or existing[4],
                            "address": lead.get("address") or existing[5],
                            "phone": lead.get("phone") or existing[6],
                            "website": lead.get("website") or existing[7],
                            "city": lead.get("city"), "region": lead.get("region"), "country": lead.get("country"),
                            "latitude": lead.get("latitude"), "longitude": lead.get("longitude"), "email": lead.get("email"),
                            "instagram": lead.get("instagram"), "facebook": lead.get("facebook"), "twitter": lead.get("twitter"),
                            "linkedin": lead.get("linkedin"), "youtube": lead.get("youtube"), "tiktok": lead.get("tiktok"),
                            "social_evidence": json.dumps(lead.get("social_evidence") or []), "social_status": lead.get("social_status"),
                            "category": lead.get("category"), "rating": lead.get("rating"), "reviews": lead.get("reviews"),
                            "maps_url": lead.get("maps_url"), "source": ", ".join(merged_sources), "sources": json.dumps(merged_sources),
                            "lead_score": lead.get("lead_score"), "lead_type": lead.get("lead_type"),
                            "ai_recommendation": lead.get("ai_recommendation"), "ai_reason": lead.get("ai_reason"),
                            "enrichment_source": lead.get("enrichment_source"), "enrichment_at": lead.get("enrichment_at"),
                        }
                        row = None
                        if existing[1] == current_user["id"]:
                            cur.execute("""UPDATE leads SET provider_place_id=%(provider_place_id)s,name=%(name)s,address=%(address)s,city=%(city)s,region=%(region)s,country=%(country)s,
                                latitude=%(latitude)s,longitude=%(longitude)s,phone=%(phone)s,website=%(website)s,email=%(email)s,instagram=COALESCE(NULLIF(%(instagram)s,''),instagram),facebook=COALESCE(NULLIF(%(facebook)s,''),facebook),twitter=COALESCE(NULLIF(%(twitter)s,''),twitter),linkedin=COALESCE(NULLIF(%(linkedin)s,''),linkedin),youtube=COALESCE(NULLIF(%(youtube)s,''),youtube),tiktok=COALESCE(NULLIF(%(tiktok)s,''),tiktok),social_evidence=CASE WHEN %(social_evidence)s::jsonb <> '[]'::jsonb THEN %(social_evidence)s::jsonb ELSE social_evidence END,social_status=COALESCE(NULLIF(%(social_status)s,''),social_status),
                                category=%(category)s,rating=%(rating)s,reviews=%(reviews)s,maps_url=%(maps_url)s,source=%(source)s,sources=%(sources)s::jsonb,lead_score=%(lead_score)s,
                                lead_type=%(lead_type)s,ai_recommendation=%(ai_recommendation)s,ai_reason=%(ai_reason)s,enrichment_source=%(enrichment_source)s,enrichment_at=%(enrichment_at)s,updated_at=NOW()
                                WHERE id=%(id)s RETURNING id,status,notes,created_at,updated_at,assigned_to,last_contacted_date,next_followup_date,is_duplicate,duplicate_of""",
                                {**updates, "id": existing[0]})
                            row = cur.fetchone()
                    else:
                        cur.execute("""INSERT INTO leads(
                            user_id,place_id,provider_place_id,name,address,city,region,country,latitude,longitude,
                            phone,website,email,instagram,facebook,twitter,linkedin,youtube,tiktok,social_evidence,social_status,category,rating,reviews,maps_url,source,sources,
                            lead_score,lead_type,ai_recommendation,ai_reason,status,assigned_to,enrichment_source,enrichment_at,created_at,updated_at)
                            VALUES(%(user_id)s,%(place_id)s,%(provider_place_id)s,%(name)s,%(address)s,%(city)s,%(region)s,%(country)s,%(latitude)s,%(longitude)s,
                            %(phone)s,%(website)s,%(email)s,%(instagram)s,%(facebook)s,%(twitter)s,%(linkedin)s,%(youtube)s,%(tiktok)s,%(social_evidence)s::jsonb,%(social_status)s,%(category)s,%(rating)s,%(reviews)s,%(maps_url)s,%(source)s,%(sources)s::jsonb,
                            %(lead_score)s,%(lead_type)s,%(ai_recommendation)s,%(ai_reason)s,'New',%(assigned_to)s,%(enrichment_source)s,%(enrichment_at)s,NOW(),NOW())
                            RETURNING id,status,notes,created_at,updated_at,assigned_to,last_contacted_date,next_followup_date,is_duplicate,duplicate_of""", {
                                **lead, "user_id": current_user["id"], "assigned_to": current_user["id"], "sources": json.dumps(lead.get("sources") or []), "social_evidence": json.dumps(lead.get("social_evidence") or []),
                                "enrichment_source": lead.get("enrichment_source"), "enrichment_at": lead.get("enrichment_at")})
                        row = cur.fetchone()
                    if existing:
                        lead["is_duplicate"] = True
                        lead["duplicate_of"] = existing[0]
                        duplicates.append({"id": existing[0], "name": existing[4], "status": existing[9], "assigned_to": existing[13]})
                        log_activity(cur,current_user["id"],existing[0],"duplicate_skipped",{"matched_lead_id":existing[0],"matched_name":existing[4]})
                        # Existing CRM record is updated only when owned by this user; the repeated search result is never re-added.
                        continue
                    assigned_value = row[5] if row else current_user["id"]
                    lead.update({"id": row[0], "status": row[1], "notes": row[2], "created_at": str(row[3]), "updated_at": str(row[4]), "assigned_to": assigned_value})
                    saved.append(lead)

                cur.execute("""INSERT INTO search_history(user_id,city,category,keyword,city_wide,result_count,sources_used)
                               VALUES(%s,%s,%s,%s,%s,%s,%s::jsonb)""",
                            (current_user["id"], location["resolved_city"], category, keyword, req.city_wide, len(saved), json.dumps(used_sources)))
                log_activity(cur, current_user["id"], None, "lead_search", {"city": city, "category": category, "keyword": keyword, "result_count": len(saved)})
            conn.commit()
    else:
        for lead in leads:
            lead.update({"id": None, "status": "New", "notes": None, "created_at": None, "updated_at": None, "assigned_to": current_user["id"]})
            saved.append(lead)

    saved.sort(key=lambda x: (0 if x.get("lead_type")=="Hot Lead" else 1 if x.get("lead_type")=="Warm Lead" else 2, -(x.get("lead_score") or 0)))
    return {"city_resolved": location["resolved_city"], "category": category, "keyword": keyword,
            "city_wide": req.city_wide, "radius_meters": radius_meters, "count": len(saved),
            "duplicates_skipped": len(duplicates), "duplicates": duplicates,
            "sources_used": used_sources, "warnings": list(dict.fromkeys(warnings)), "leads": saved}
# Public social-profile discovery using the existing SerpApi Google Search engine.
# Search results are treated as candidates; business-name/location checks reduce false matches.
_SOCIAL_HOSTS = {
    "linkedin": ("linkedin.com/company/", "linkedin.com/showcase/"),
    "instagram": ("instagram.com/",),
    "facebook": ("facebook.com/", "fb.com/"),
    "youtube": ("youtube.com/", "youtu.be/"),
    "twitter": ("twitter.com/", "x.com/"),
    "tiktok": ("tiktok.com/",),
}

def discover_public_social_profiles(lead: dict):
    """Find candidate business profiles via Google results; never infer absence from no results."""
    if not SERPAPI_API_KEY or not lead or not lead.get("name"):
        return {"social_status": "not_verified", "social_evidence": []}
    name = str(lead.get("name") or "").strip()
    city = str(lead.get("city") or lead.get("address") or "").strip()
    # A single bounded Google query per lead limits usage of the existing SerpApi plan.
    q = f'"{name}" {city} (site:linkedin.com/company OR site:instagram.com OR site:facebook.com OR site:youtube.com OR site:x.com OR site:tiktok.com)'
    try:
        data = serpapi_get({"engine": "google", "q": q, "hl": "en", "gl": "in", "num": 10}, timeout=18)
    except Exception:
        return {"social_status": "not_verified", "social_evidence": []}
    profiles = {}
    evidence = []
    name_tokens = [x.casefold() for x in re.findall(r"[a-zA-Z0-9]+", name) if len(x) > 2]
    for item in (data.get("organic_results") or [])[:10]:
        link = str(item.get("link") or "").strip()
        title = str(item.get("title") or "")
        snippet = str(item.get("snippet") or "")
        if not link or not _ai_safe_public_url(link):
            continue
        low = link.casefold()
        platform = next((key for key, hosts in _SOCIAL_HOSTS.items() if any(host in low for host in hosts)), None)
        if not platform:
            continue
        # Require a meaningful business-name token match in the result title/snippet.
        haystack = (title + " " + snippet).casefold()
        matched = sum(1 for token in name_tokens if token in haystack)
        if not name_tokens or matched < min(2, len(name_tokens)):
            continue
        # Exclude generic social landing/search URLs, posts, and login pages where possible.
        path = urlparse(link).path.strip("/").casefold()
        if not path or path in {"login", "share", "search", "watch", "explore"}:
            continue
        if platform in profiles:
            continue
        profiles[platform] = link
        evidence.append({"platform": platform, "url": link, "title": title[:240], "snippet": snippet[:500], "source": "SerpApi Google Search", "verified": True})
    return {**profiles, "social_status": "profiles_found" if profiles else "not_found_in_public_search", "social_evidence": evidence}

# ============================================================

# Saved leads / filters / updates

# ============================================================

LEAD_SELECT = """SELECT id,name,address,city,region,country,phone,website,email,instagram,facebook,twitter,linkedin,youtube,tiktok,social_evidence,social_status,category,rating,reviews,maps_url,source,lead_score,lead_type,ai_recommendation,ai_reason,status,place_id,created_at,updated_at,notes,latitude,longitude,provider_place_id,sources,enrichment_source,enrichment_at,assigned_to,last_contacted_date,next_followup_date,is_duplicate,duplicate_of FROM leads"""

def row_to_lead(row):

    return {

        "id": row[0], "name": row[1], "address": row[2], "city": row[3], "region": row[4], "country": row[5],

        "phone": row[6], "website": row[7], "email": row[8], "instagram": row[9], "facebook": row[10], "twitter": row[11],
        "linkedin": row[12], "youtube": row[13], "tiktok": row[14], "social_evidence": row[15] or [], "social_status": row[16],

        "category": row[17], "rating": row[18], "reviews": row[19], "maps_url": row[20], "source": row[21], "lead_score": row[22],

        "lead_type": row[23], "ai_recommendation": row[24], "ai_reason": row[25], "status": row[26], "place_id": row[27],

        "created_at": str(row[28]), "updated_at": str(row[29]), "notes": row[30], "latitude": row[31], "longitude": row[32], "provider_place_id": row[33],

        "sources": row[34] or [], "enrichment_source": row[35], "enrichment_at": str(row[36]) if row[36] else None,
        "whatsapp_url": whatsapp_url(row[6]), "assigned_to": row[37],
        "last_contacted_date": str(row[38]) if row[38] else None, "next_followup_date": str(row[39]) if row[39] else None,
        "is_duplicate": bool(row[40]), "duplicate_of": row[41],

    }

@app.get("/api/leads")
def get_saved_leads(has_website: Optional[bool]=None,has_phone: Optional[bool]=None,has_email: Optional[bool]=None,has_social: Optional[bool]=None,min_score: Optional[int]=Query(None,ge=0,le=100),lead_type: Optional[str]=None,min_rating: Optional[float]=Query(None,ge=0,le=5),category: Optional[str]=None,city: Optional[str]=None,source: Optional[str]=None,status: Optional[str]=None,assigned_to: Optional[int]=None,date_from: Optional[str]=None,date_to: Optional[str]=None,current_user: dict=Depends(get_current_user)):
    scope,params=lead_access_clause(current_user); clauses=[scope]; params=list(params)
    for flag,col in [(has_website,"website"),(has_phone,"phone"),(has_email,"email")]:
        if flag is True: clauses.append(f"NULLIF(TRIM({col}),'') IS NOT NULL")
        elif flag is False: clauses.append(f"NULLIF(TRIM({col}),'') IS NULL")
    if has_social is True: clauses.append("(NULLIF(TRIM(instagram),'') IS NOT NULL OR NULLIF(TRIM(facebook),'') IS NOT NULL OR NULLIF(TRIM(twitter),'') IS NOT NULL OR NULLIF(TRIM(linkedin),'') IS NOT NULL OR NULLIF(TRIM(youtube),'') IS NOT NULL OR NULLIF(TRIM(tiktok),'') IS NOT NULL)")
    elif has_social is False: clauses.append("(NULLIF(TRIM(instagram),'') IS NULL AND NULLIF(TRIM(facebook),'') IS NULL AND NULLIF(TRIM(twitter),'') IS NULL AND NULLIF(TRIM(linkedin),'') IS NULL AND NULLIF(TRIM(youtube),'') IS NULL AND NULLIF(TRIM(tiktok),'') IS NULL)")
    if min_score is not None: clauses.append("COALESCE(lead_score,0)>=%s"); params.append(min_score)
    if lead_type: clauses.append("lead_type=%s"); params.append(lead_type)
    if min_rating is not None: clauses.append("COALESCE(rating,0)>=%s"); params.append(min_rating)
    for val,col in [(category,"category"),(city,"city"),(source,"source"),(status,"status")]:
        if val: clauses.append(f"{col} ILIKE %s"); params.append(f"%{val}%")
    start_utc, end_utc = local_date_range_utc(date_from, date_to)
    if start_utc is not None:
        clauses.append("created_at >= %s"); params.append(start_utc)
    if end_utc is not None:
        clauses.append("created_at < %s"); params.append(end_utc)
    if assigned_to is not None:
        if is_team_leader(current_user):
            if assigned_to != current_user["id"]:
                with db_conn() as conn:
                    with conn.cursor() as cur:
                        cur.execute("SELECT id FROM users WHERE id=%s AND team_leader_id=%s",(assigned_to,current_user["id"]))
                        if not cur.fetchone(): raise HTTPException(status_code=403,detail="Agent is outside your team.")
        elif not is_admin(current_user) and assigned_to!=current_user["id"]:
            raise HTTPException(status_code=403,detail="Agents can only view their assigned leads.")
        clauses.append("assigned_to=%s"); params.append(assigned_to)
    order="CASE WHEN lead_type='Hot Lead' THEN 0 WHEN lead_type='Warm Lead' THEN 1 ELSE 2 END,COALESCE(lead_score,0) DESC,created_at DESC"
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(LEAD_SELECT+" WHERE "+" AND ".join(clauses)+" ORDER BY "+order,params); rows=cur.fetchall()
    return {"success":True,"total":len(rows),"leads":[row_to_lead(r) for r in rows]}

@app.get("/api/admin/agents")
def admin_list_agents(current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""SELECT id,email,name,COALESCE(role,'Agent'),COALESCE(active,TRUE),created_at,COALESCE(approval_status,'Approved'),phone,COALESCE(auth_provider,'password')
                           FROM users ORDER BY CASE WHEN role='Admin' THEN 0 ELSE 1 END,name NULLS LAST,email""")
            rows=cur.fetchall()
    return {"agents":[{"id":r[0],"email":r[1],"name":r[2],"role":r[3],"active":r[4],"created_at":str(r[5]) if r[5] else None,"approval_status":r[6],"phone":r[7],"auth_provider":r[8]} for r in rows]}

@app.post("/api/admin/agents")
def admin_create_agent(req:AdminAgentCreateRequest,current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    email=req.email.strip().lower()
    name=(req.name or "").strip()
    if "@" not in email or "." not in email.split("@")[-1]:
        raise HTTPException(status_code=400,detail="Please enter a valid email address.")
    password_hash,salt=hash_password(req.password)
    try:
        with db_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT id FROM users WHERE lower(email)=lower(%s)",(email,))
                if cur.fetchone():
                    raise HTTPException(status_code=400,detail="An account with this email already exists.")
                cur.execute("""INSERT INTO users(email,name,password_hash,salt,role,team_leader_id,active,approval_status,auth_provider)
                               VALUES(%s,%s,%s,%s,'Agent',NULL,TRUE,'Approved','password') RETURNING id""",
                            (email,name or None,password_hash,salt))
                agent_id=cur.fetchone()[0]
                log_activity(cur,current_user["id"],None,"agent_created",{"agent_id":agent_id,"email":email})
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500,detail=f"Agent creation failed: {type(exc).__name__}.")
    return {"success":True,"agent":{"id":agent_id,"email":email,"name":name,"role":"Agent","active":True}}

@app.patch("/api/admin/agents/{user_id}")
def admin_update_agent(user_id:int,req:AdminAgentUpdateRequest,current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    data=req.model_dump(exclude_unset=True)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,email,COALESCE(role,'Agent'),COALESCE(active,TRUE) FROM users WHERE id=%s",(user_id,))
            target=cur.fetchone()
            if not target: raise HTTPException(status_code=404,detail="Account not found.")
            if target[2]=="Admin" and data.get("active") is False:
                raise HTTPException(status_code=400,detail="The Admin account cannot be disabled.")
            if user_id==current_user["id"] and data.get("active") is False:
                raise HTTPException(status_code=400,detail="You cannot disable your own Admin account.")
            sets=[]; vals=[]
            if "email" in data and data["email"] is not None:
                email=str(data["email"]).strip().lower()
                if "@" not in email or "." not in email.split("@")[-1]:
                    raise HTTPException(status_code=400,detail="Invalid email address.")
                sets.append("email=%s"); vals.append(email)
            if "name" in data:
                sets.append("name=%s"); vals.append((data["name"] or "").strip() or None)
            if data.get("password"):
                ph,salt=hash_password(data["password"])
                sets.extend(["password_hash=%s","salt=%s"]); vals.extend([ph,salt])
            if "active" in data and data["active"] is not None:
                sets.append("active=%s"); vals.append(bool(data["active"]))
                if data["active"]:
                    sets.append("approval_status=%s"); vals.append("Approved")
            if not sets:
                raise HTTPException(status_code=400,detail="No changes supplied.")
            vals.append(user_id)
            try:
                cur.execute("UPDATE users SET "+",".join(sets)+" WHERE id=%s",vals)
            except psycopg.errors.UniqueViolation:
                raise HTTPException(status_code=400,detail="An account with that email already exists.")
            log_activity(cur,current_user["id"],None,"agent_updated",{"user_id":user_id,"fields":list(data.keys())})
        conn.commit()
    return {"success":True,"user_id":user_id}


@app.post("/api/admin/agents/{user_id}/approve")
def admin_approve_agent(user_id:int,current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,COALESCE(role,'Agent'),COALESCE(active,TRUE) FROM users WHERE id=%s",(user_id,))
            target=cur.fetchone()
            if not target: raise HTTPException(status_code=404,detail="Account not found.")
            if target[1]=='Admin': raise HTTPException(status_code=400,detail="The Admin account is already approved.")
            cur.execute("UPDATE users SET approval_status='Approved',active=TRUE WHERE id=%s",(user_id,))
            log_activity(cur,current_user['id'],None,'agent_approved',{'user_id':user_id})
        conn.commit()
    return {"success":True,"user_id":user_id,"approval_status":"Approved","active":True}

@app.post("/api/admin/agents/{user_id}/reset-password")
def admin_reset_agent_password(user_id:int, req:AdminResetPasswordRequest, current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    password_hash, salt = hash_password(req.new_password)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,email,COALESCE(role,'Agent'),COALESCE(active,TRUE) FROM users WHERE id=%s", (user_id,))
            target = cur.fetchone()
            if not target:
                raise HTTPException(status_code=404, detail="Account not found.")
            if target[2] == "Admin" and user_id != current_user["id"]:
                raise HTTPException(status_code=403, detail="Admin password can only be changed from the Admin account.")
            if not target[3]:
                raise HTTPException(status_code=400, detail="This account is disabled.")
            cur.execute("UPDATE users SET password_hash=%s,salt=%s WHERE id=%s", (password_hash, salt, user_id))
            log_activity(cur, current_user["id"], None, "password_reset", {"user_id": user_id, "email": target[1]})
        conn.commit()
    return {"success":True,"user_id":user_id}

@app.delete("/api/admin/agents/{user_id}")
def admin_disable_agent(user_id:int,current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,COALESCE(role,'Agent') FROM users WHERE id=%s",(user_id,))
            target=cur.fetchone()
            if not target: raise HTTPException(status_code=404,detail="Account not found.")
            if target[1]=="Admin" or user_id==current_user["id"]:
                raise HTTPException(status_code=400,detail="The Admin account cannot be disabled or deleted.")
            cur.execute("UPDATE users SET active=FALSE WHERE id=%s",(user_id,))
            log_activity(cur,current_user["id"],None,"agent_disabled",{"user_id":user_id})
        conn.commit()
    return {"success":True,"user_id":user_id,"active":False}

@app.get("/api/team-members")
def team_members(current_user: dict=Depends(get_current_user)):
    scope,params=accessible_user_clause(current_user,"u")
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT u.id,u.email,u.name,COALESCE(u.role,'Agent'),u.team_leader_id,t.name,COALESCE(u.active,TRUE),COALESCE(u.approval_status,'Approved'),COALESCE(u.auth_provider,'password'),u.phone FROM users u LEFT JOIN users t ON t.id=u.team_leader_id WHERE "+scope+" ORDER BY u.name NULLS LAST,u.email",params)
            rows=cur.fetchall()
    return {"team_members":[{"id":r[0],"email":r[1],"name":r[2],"role":r[3],"team_leader_id":r[4],"team_leader_name":r[5],"active":r[6],"approval_status":r[7],"auth_provider":r[8],"phone":r[9]} for r in rows]}

@app.patch("/api/team-members/{user_id}/role")
def update_team_member_role(user_id:int, role:str, current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    role=role.strip().title()
    if role not in VALID_ROLES:
        raise HTTPException(status_code=400, detail="Role must be Admin or Agent.")
    if role == "Admin":
        raise HTTPException(status_code=400, detail="There must be exactly one separate Admin account.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,COALESCE(role,'Agent') FROM users WHERE id=%s",(user_id,))
            target=cur.fetchone()
            if not target: raise HTTPException(status_code=404, detail="Team member not found.")
            if target[1]=='Admin' and role!='Admin':
                raise HTTPException(status_code=400, detail="The single Admin account cannot be changed to Agent.")
            if user_id==current_user["id"] and role!="Admin":
                raise HTTPException(status_code=400, detail="The Admin account cannot remove its own Admin role.")
            cur.execute("UPDATE users SET role=%s,team_leader_id=NULL WHERE id=%s",(role,user_id))
            log_activity(cur,current_user['id'],None,'team_role_changed',{'user_id':user_id,'role':role})
        conn.commit()
    return {"success":True,"user_id":user_id,"role":role}

@app.get("/api/dashboard/followups")
def dashboard_followups(current_user: dict=Depends(get_current_user)):
    scope,base_params=lead_access_clause(current_user,"l"); now=datetime.now(timezone.utc); start_utc,end_utc,today_name=local_today_utc_bounds()
    with db_conn() as conn:
        with conn.cursor() as cur:
            sql=LEAD_SELECT.replace("FROM leads","FROM leads l")+" WHERE "+scope+" AND next_followup_date IS NOT NULL AND next_followup_date < %s AND status NOT IN ('Converted','Lost') ORDER BY next_followup_date ASC"
            cur.execute(sql,base_params+[now]); missed=[row_to_lead(r) for r in cur.fetchall()]
            sql=LEAD_SELECT.replace("FROM leads","FROM leads l")+" WHERE "+scope+" AND next_followup_date >= %s AND next_followup_date < %s AND status NOT IN ('Converted','Lost') ORDER BY next_followup_date ASC"
            cur.execute(sql,base_params+[start_utc,end_utc]); today_rows=[row_to_lead(r) for r in cur.fetchall()]
    return {"today":today_rows,"missed":missed,"today_count":len(today_rows),"missed_count":len(missed),"date":today_name,"timezone":user_timezone_name()}

@app.get("/api/dashboard/summary")
def dashboard_summary(current_user:dict=Depends(get_current_user)):
    scope,params=lead_access_clause(current_user,'l'); start_utc,end_utc,_=local_today_utc_bounds(); now=datetime.now(timezone.utc)
    with db_conn() as conn:
        with conn.cursor() as cur:
            # Keep the access-scope parameters separate from the date parameters so PostgreSQL
            # never binds a date value into the lead-access placeholders.
            sql=f"SELECT COUNT(*),COUNT(*) FILTER(WHERE l.last_contacted_date IS NOT NULL),COUNT(*) FILTER(WHERE l.status='Converted'),COUNT(*) FILTER(WHERE l.next_followup_date IS NOT NULL AND l.next_followup_date >= %s AND l.next_followup_date < %s AND l.status NOT IN ('Converted','Lost')),COUNT(*) FILTER(WHERE l.next_followup_date IS NOT NULL AND l.next_followup_date < %s AND l.status NOT IN ('Converted','Lost')) FROM leads l WHERE {scope}"
            cur.execute(sql,[start_utc,end_utc,now,*params]); total,contacted,success,today,missed=cur.fetchone()
            cur.execute('SELECT COUNT(*) FROM call_logs c JOIN leads l ON l.id=c.lead_id WHERE '+scope,params); calls=cur.fetchone()[0] or 0
    return {'total_leads':total or 0,'contacted_leads':contacted or 0,'success_leads':success or 0,'total_calls':calls,'today_followups':today or 0,'missed_followups':missed or 0}

@app.get("/api/dashboard/agent-performance")
def agent_performance(date_from:Optional[str]=None,date_to:Optional[str]=None,current_user:dict=Depends(get_current_user)):
    start=end=None
    if date_from:
        try: start=datetime.fromisoformat(date_from).replace(tzinfo=timezone.utc)
        except ValueError: raise HTTPException(status_code=400,detail="Invalid date_from.")
    if date_to:
        try: end=datetime.fromisoformat(date_to).replace(tzinfo=timezone.utc)+timedelta(days=1)
        except ValueError: raise HTTPException(status_code=400,detail="Invalid date_to.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            if is_admin(current_user):
                cur.execute("SELECT u.id,u.name,u.email,COALESCE(u.role,'Agent') FROM users u WHERE COALESCE(u.active,TRUE)=TRUE ORDER BY u.name NULLS LAST,u.id")
                users=cur.fetchall()
            else:
                cur.execute("SELECT u.id,u.name,u.email,COALESCE(u.role,'Agent') FROM users u WHERE u.id=%s", (current_user["id"],))
                own=cur.fetchone()
                cur.execute("SELECT u.id,u.name,u.email,COALESCE(u.role,'Agent') FROM users u WHERE COALESCE(u.active,TRUE)=TRUE AND COALESCE(u.role,'Agent')='Agent' AND u.id<>%s ORDER BY u.name NULLS LAST,u.id", (current_user["id"],))
                others=cur.fetchall()
                cur.execute("SELECT u.id,u.name,u.email,COALESCE(u.role,'Agent') FROM users u WHERE COALESCE(u.active,TRUE)=TRUE AND COALESCE(u.role,'Agent')='Admin' ORDER BY u.id LIMIT 1")
                admins=cur.fetchall()
                users=list(admins + ([own] if own else []) + list(others))
            out=[]
            for uid,name,email,role in users:
                lw=['assigned_to=%s']; lp=[uid]
                if start: lw.append('created_at >= %s'); lp.append(start)
                if end: lw.append('created_at < %s'); lp.append(end)
                cur.execute('SELECT COUNT(*) FROM leads WHERE '+' AND '.join(lw),lp); assigned=cur.fetchone()[0] or 0
                cw=['called_by=%s']; cp=[uid]
                if start: cw.append('call_date >= %s'); cp.append(start)
                if end: cw.append('call_date < %s'); cp.append(end)
                cur.execute("SELECT COUNT(*),COUNT(*) FILTER(WHERE lower(trim(call_outcome)) IN ('connected','interested','follow-up','converted')),COUNT(*) FILTER(WHERE lower(trim(call_outcome))='interested'),COUNT(*) FILTER(WHERE lower(trim(call_outcome))='follow-up'),COUNT(*) FILTER(WHERE lower(trim(call_outcome))='converted'),COUNT(DISTINCT lead_id) FILTER(WHERE lower(trim(call_outcome)) IN ('connected','interested','follow-up','converted')),COUNT(*) FILTER(WHERE lower(trim(call_outcome)) IN ('not interested','no answer','failed','failure')) FROM call_logs WHERE "+' AND '.join(cw),cp)
                calls,connected,interested,followups,converted_calls,contacts,failures=[x or 0 for x in cur.fetchone()]
                success=interested+converted_calls
                decided=success+failures
                success_rate=round(success*100/decided,1) if decided else 0.0
                contact_rate=round(contacts*100/assigned,1) if assigned else 0.0
                followup_rate=round(followups*100/assigned,1) if assigned else 0.0
                # Balanced score: contact activity, connection quality, follow-up discipline and successful outcomes.
                score=round(min(100,(contact_rate*0.25)+(min(100,connected*100/max(calls,1))*0.20)+(min(100,followup_rate)*0.20)+(success_rate*0.35)),1)
                out.append({'id':uid,'name':name,'email':email,'role':role,'assigned_leads':assigned,'calls':calls,'contacts':contacts,'connected':connected,'interested':interested,'followups':followups,'converted':converted_calls,'success':success,'failure':failures,'contact_rate':contact_rate,'conversion_rate':success_rate,'score':score})
    out.sort(key=lambda x:(-x['score'],-x['converted'],-x['contacts'],-x['calls'],x['name'] or ''))
    for i,row in enumerate(out,1): row['rank']=i
    if not is_admin(current_user):
        admin_rows=[r for r in out if r['role']=='Admin']
        own_row=next((r for r in out if r['id']==current_user['id']), None)
        top5=[r for r in out if r['role']=='Agent' and r['id']!=current_user['id']][:5]
        selected=admin_rows + top5 + ([own_row] if own_row else [])
        for r in selected:
            if r['role']=='Admin': r['rank']=0
        out=sorted(selected, key=lambda x:(0 if x['role']=='Admin' else 1, x['rank']))
    return {'performance':out,'period':{'date_from':date_from,'date_to':date_to}}

@app.patch("/api/leads/{lead_id}/status")
def update_status(lead_id:int,req:StatusUpdateRequest,current_user:dict=Depends(get_current_user)):
    if req.status not in VALID_STATUSES: raise HTTPException(status_code=400,detail="Invalid status.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user,True); cur.execute("UPDATE leads SET status=%s,updated_at=NOW() WHERE id=%s",(req.status,lead_id)); log_activity(cur,current_user["id"],lead_id,"status_changed",{"status":req.status})
        conn.commit()
    return {"success":True,"lead_id":lead_id,"status":req.status}

@app.patch("/api/leads/{lead_id}/assignment")
def assign_lead(lead_id:int,req:AssignmentRequest,current_user:dict=Depends(get_current_user)):
    if not is_admin(current_user): raise HTTPException(status_code=403,detail="Only Admin can assign leads.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user,True)
            if req.assigned_to is not None:
                cur.execute("SELECT id FROM users WHERE id=%s AND COALESCE(active,TRUE)=TRUE",(req.assigned_to,))
                if not cur.fetchone(): raise HTTPException(status_code=404,detail="Team member not found or outside your team.")
            cur.execute("SELECT name FROM leads WHERE id=%s", (lead_id,))
            lead_row = cur.fetchone()
            cur.execute("UPDATE leads SET assigned_to=%s,updated_at=NOW() WHERE id=%s",(req.assigned_to,lead_id)); log_activity(cur,current_user["id"],lead_id,"assigned",{"assigned_to":req.assigned_to})
            if req.assigned_to is not None:
                cur.execute("INSERT INTO notifications(recipient_id,actor_id,lead_id,notification_type,title,message) VALUES(%s,%s,%s,%s,%s,%s)",
                            (req.assigned_to, current_user["id"], lead_id, 'assignment', 'New lead assigned', f'Admin assigned {lead_row[0] if lead_row else "a lead"} to you.'))
        conn.commit()
    return {"success":True,"lead_id":lead_id,"assigned_to":req.assigned_to}

@app.patch("/api/leads/{lead_id}/followup")
def set_followup(lead_id:int,req:FollowupRequest,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user,True); cur.execute("UPDATE leads SET next_followup_date=%s,updated_at=NOW() WHERE id=%s",(req.next_followup_date,lead_id)); log_activity(cur,current_user["id"],lead_id,"followup_changed",{"next_followup_date":req.next_followup_date.isoformat() if req.next_followup_date else None})
        conn.commit()
    return {"success":True,"lead_id":lead_id,"next_followup_date":req.next_followup_date}

@app.post("/api/leads/{lead_id}/calls")
def add_call_log(lead_id:int,req:CallLogRequest,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user,True); call_date=req.call_date or datetime.now(timezone.utc)
            cur.execute("INSERT INTO call_logs(lead_id,called_by,call_date,call_outcome,notes) VALUES(%s,%s,%s,%s,%s) RETURNING id",(lead_id,current_user["id"],call_date,req.call_outcome.strip(),req.notes)); call_id=cur.fetchone()[0]
            cur.execute("UPDATE leads SET last_contacted_date=%s,status=CASE WHEN status='New' THEN 'Contacted' ELSE status END,updated_at=NOW() WHERE id=%s",(call_date,lead_id)); log_activity(cur,current_user["id"],lead_id,"call_logged",{"call_id":call_id,"outcome":req.call_outcome.strip()})
        conn.commit()
    return {"success":True,"call_id":call_id,"lead_id":lead_id,"last_contacted_date":str(call_date)}

@app.get("/api/leads/{lead_id}/calls")
def get_call_logs(lead_id:int,current_user:dict=Depends(get_current_user)):
    # Always authorize the lead before reading its call history.
    # This used to call lead_owner_check without current_user, which caused:
    # TypeError: lead_owner_check() missing 1 required positional argument: 'user'
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur, lead_id, current_user)
            cur.execute(
                """
                SELECT c.id, c.called_by, c.call_date, c.call_outcome, c.notes,
                       u.name, u.email
                FROM call_logs c
                LEFT JOIN users u ON u.id = c.called_by
                WHERE c.lead_id = %s
                ORDER BY c.call_date DESC, c.id DESC
                """,
                (lead_id,),
            )
            rows = cur.fetchall()
    return {
        "calls": [
            {
                "id": r[0],
                "called_by": r[1],
                "call_date": str(r[2]),
                "call_outcome": r[3],
                "notes": r[4],
                "called_by_name": r[5],
                "called_by_email": r[6],
            }
            for r in rows
        ]
    }

@app.patch("/api/leads/{lead_id}/notes")
def update_notes(lead_id:int,req:NotesUpdateRequest,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id, current_user,True); cur.execute("UPDATE leads SET notes=%s,updated_at=NOW() WHERE id=%s",((req.notes or "")[:5000],lead_id)); log_activity(cur,current_user["id"],lead_id,"notes_changed",{})
        conn.commit()
    return {"success":True,"lead_id":lead_id,"notes":req.notes or ""}

@app.delete("/api/leads/{lead_id}")
def delete_lead(lead_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user,True); log_activity(cur,current_user["id"],lead_id,"delete_requested",{}); cur.execute("DELETE FROM leads WHERE id=%s",(lead_id,))
        conn.commit()
    return {"success":True}

@app.patch("/api/leads/{lead_id}")
def edit_lead(lead_id:int,req:LeadUpdateRequest,current_user:dict=Depends(get_current_user)):
    data=req.model_dump(exclude_unset=True)
    if "status" in data and data["status"] not in VALID_STATUSES: raise HTTPException(status_code=400,detail="Invalid status.")
    if "assigned_to" in data and not is_admin(current_user): raise HTTPException(status_code=403,detail="Only Admin can assign leads.")
    allowed={"name","address","city","region","country","phone","website","email","instagram","facebook","twitter","category","notes","status","assigned_to","last_contacted_date","next_followup_date"}; data={k:v for k,v in data.items() if k in allowed}
    if not data: raise HTTPException(status_code=400,detail="No changes supplied.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user,True)
            if data.get("assigned_to") is not None:
                cur.execute("SELECT id FROM users WHERE id=%s",(data["assigned_to"],));
                if not cur.fetchone(): raise HTTPException(status_code=404,detail="Team member not found.")
            cur.execute(f"UPDATE leads SET {','.join(f'{k}=%s' for k in data)},updated_at=NOW() WHERE id=%s",list(data.values())+[lead_id]); log_activity(cur,current_user["id"],lead_id,"lead_updated",data)
        conn.commit()
    return {"success":True,"lead_id":lead_id}

@app.post("/api/leads/{lead_id}/whatsapp")
def lead_whatsapp(lead_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            lead_owner_check(cur,lead_id,current_user); cur.execute("SELECT phone FROM leads WHERE id=%s",(lead_id,)); phone=cur.fetchone()[0]
    url=whatsapp_url(phone)
    if not url: raise HTTPException(status_code=400,detail="Lead has no usable phone number.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            log_activity(cur,current_user["id"],lead_id,"whatsapp_opened",{})
        conn.commit()
    return {"success":True,"whatsapp_url":url,"manual_only":True}

@app.post("/api/leads/bulk")
def bulk_lead_action(req:BulkLeadActionRequest,current_user:dict=Depends(get_current_user)):
    ids=list(dict.fromkeys(req.lead_ids)); action=req.action.strip().lower()
    with db_conn() as conn:
        with conn.cursor() as cur:
            scope,params=lead_access_clause(current_user)
            if action=="status":
                if req.status not in VALID_STATUSES: raise HTTPException(status_code=400,detail="Invalid status.")
                cur.execute(f"UPDATE leads SET status=%s,updated_at=NOW() WHERE id=ANY(%s) AND {scope}",(req.status,ids,*params)); count=cur.rowcount
            elif action=="assign":
                require_admin(current_user)
                if req.assigned_to is None: raise HTTPException(status_code=400,detail="assigned_to is required.")
                cur.execute("SELECT id FROM users WHERE id=%s AND COALESCE(active,TRUE)=TRUE AND COALESCE(role,'Agent')='Agent'",(req.assigned_to,));
                if not cur.fetchone(): raise HTTPException(status_code=404,detail="Team member not found.")
                cur.execute("SELECT id,name FROM leads WHERE id=ANY(%s)",(ids,)); assign_rows=cur.fetchall()
                cur.execute("UPDATE leads SET assigned_to=%s,updated_at=NOW() WHERE id=ANY(%s)",(req.assigned_to,ids)); count=cur.rowcount
            elif action=="delete":
                cur.execute(f"SELECT id FROM leads WHERE id=ANY(%s) AND {scope}",(ids,*params)); delete_rows=cur.fetchall()
                for (lead_id,) in delete_rows: log_activity(cur,current_user["id"],lead_id,"bulk_delete_requested",{})
                cur.execute(f"DELETE FROM leads WHERE id=ANY(%s) AND {scope}",(ids,*params)); count=cur.rowcount
            else: raise HTTPException(status_code=400,detail="Unsupported bulk action.")
            if action == "status":
                cur.execute(f"SELECT id FROM leads WHERE id=ANY(%s) AND {scope}",(ids,*params))
                for (lead_id,) in cur.fetchall():
                    log_activity(cur,current_user["id"],lead_id,"bulk_status_changed",{"status":req.status})
            elif action == "assign":
                for lead_id,lead_name in assign_rows:
                    log_activity(cur,current_user["id"],lead_id,"bulk_assigned",{"assigned_to":req.assigned_to})
                    cur.execute("INSERT INTO notifications(recipient_id,actor_id,lead_id,notification_type,title,message) VALUES(%s,%s,%s,%s,%s,%s)",
                                (req.assigned_to, current_user["id"], lead_id, 'assignment', 'New lead assigned', f'Admin assigned {lead_name or "a lead"} to you.'))
            elif action == "delete":
                # Deletions cascade activity rows by design; record the action is not possible after the lead is deleted.
                pass
        conn.commit()
    return {"success":True,"updated":count} if action!="delete" else {"success":True,"deleted":count}

@app.get("/api/activity-log")
def activity_log(limit:int=Query(100,ge=1,le=500),lead_id:Optional[int]=None,current_user:dict=Depends(get_current_user)):
    scope,params=lead_access_clause(current_user,"l"); where=[scope]; params=list(params)
    if lead_id is not None: where.append("a.lead_id=%s"); params.append(lead_id)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT a.id,a.user_id,a.lead_id,a.action,a.details,a.created_at,u.name,u.email,l.name FROM activity_logs a LEFT JOIN users u ON u.id=a.user_id LEFT JOIN leads l ON l.id=a.lead_id WHERE "+" AND ".join(where)+" ORDER BY a.created_at DESC LIMIT %s",params+[limit]); rows=cur.fetchall()
    return {"activity":[{"id":r[0],"user_id":r[1],"lead_id":r[2],"action":r[3],"details":r[4] or {},"created_at":str(r[5]),"user_name":r[6],"user_email":r[7],"lead_name":r[8]} for r in rows]}

# Notifications

# ============================================================

@app.get("/api/notifications")
def get_notifications(limit:int=Query(100,ge=1,le=200), current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""SELECT n.id,n.actor_id,n.lead_id,n.notification_type,n.title,n.message,n.is_read,n.created_at,u.name,u.email,l.name
                       FROM notifications n LEFT JOIN users u ON u.id=n.actor_id LEFT JOIN leads l ON l.id=n.lead_id
                       WHERE n.recipient_id=%s ORDER BY n.created_at DESC,n.id DESC LIMIT %s""", (current_user["id"],limit))
            rows=cur.fetchall()
    return {"notifications":[{"id":r[0],"actor_id":r[1],"lead_id":r[2],"type":r[3],"title":r[4],"message":r[5],"is_read":bool(r[6]),"created_at":str(r[7]),"actor_name":r[8],"actor_email":r[9],"lead_name":r[10]} for r in rows], "unread_count":sum(1 for r in rows if not r[6])}

@app.post("/api/notifications/{notification_id}/read")
def mark_notification_read(notification_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("UPDATE notifications SET is_read=TRUE WHERE id=%s AND recipient_id=%s",(notification_id,current_user["id"]))
            if cur.rowcount==0: raise HTTPException(status_code=404,detail="Notification not found.")
        conn.commit()
    return {"success":True}

@app.delete("/api/notifications/{notification_id}")
def delete_notification(notification_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM notifications WHERE id=%s AND recipient_id=%s",(notification_id,current_user["id"]))
            if cur.rowcount==0: raise HTTPException(status_code=404,detail="Notification not found.")
        conn.commit()
    return {"success":True}

# Internal Chat
# ============================================================
# Private chat is participant-only. Admins may participate in chats
# but cannot inspect, delete, or control another user's private chat.
CHAT_ONLINE_SECONDS = 45


def _chat_user_is_available(cur, user_id: int) -> bool:
    cur.execute("""SELECT id,COALESCE(role,'Agent'),COALESCE(active,TRUE),COALESCE(approval_status,'Approved')
                   FROM users WHERE id=%s""", (user_id,))
    row = cur.fetchone()
    return bool(row and row[2] and row[3] == "Approved")


def _require_chat_user(cur, user_id: int):
    if not _chat_user_is_available(cur, user_id):
        raise HTTPException(status_code=403, detail="Your account is not available for chat.")


def _chat_pair(user_a: int, user_b: int):
    if user_a == user_b:
        raise HTTPException(status_code=400, detail="You cannot start a chat with yourself.")
    return (min(user_a, user_b), max(user_a, user_b))


def _chat_conversation_row(cur, conversation_id: int):
    cur.execute("""SELECT c.id,c.user1_id,c.user2_id,c.created_at,c.last_message_at,
                          u1.name,u1.email,u1.role,u1.active,u1.approval_status,
                          u2.name,u2.email,u2.role,u2.active,u2.approval_status
                   FROM chat_conversations c
                   JOIN users u1 ON u1.id=c.user1_id
                   JOIN users u2 ON u2.id=c.user2_id
                   WHERE c.id=%s""", (conversation_id,))
    return cur.fetchone()


def _chat_is_participant(conversation, user_id: int) -> bool:
    return bool(conversation and user_id in (conversation[1], conversation[2]))


def _chat_user_payload(row):
    return {"id": row[0], "name": row[1], "email": row[2], "role": row[3],
            "active": bool(row[4]), "approval_status": row[5]}


def _chat_reactions(cur, private_ids=None, team_ids=None):
    private_ids=private_ids or []
    team_ids=team_ids or []
    result={}
    if private_ids:
        cur.execute("SELECT private_message_id,reaction,COUNT(*) FROM chat_message_reactions WHERE private_message_id=ANY(%s) GROUP BY private_message_id,reaction",(private_ids,))
        for mid,reaction,count in cur.fetchall(): result.setdefault(("p",mid),{})[reaction]=int(count)
    if team_ids:
        cur.execute("SELECT team_message_id,reaction,COUNT(*) FROM chat_message_reactions WHERE team_message_id=ANY(%s) GROUP BY team_message_id,reaction",(team_ids,))
        for mid,reaction,count in cur.fetchall(): result.setdefault(("t",mid),{})[reaction]=int(count)
    return result


def _chat_message_payload(r, current_user_id, reactions):
    return {
        "id":r[0],"conversation_id":r[1],"sender_id":r[2],"body":r[3],
        "is_read":bool(r[4]),"created_at":str(r[5]),"edited_at":str(r[6]) if r[6] else None,
        "deleted_at":str(r[7]) if r[7] else None,"is_deleted":bool(r[7]),
        "reply_to_id":r[8],"reply_body":r[9],"reply_sender_id":r[10],"reply_sender_name":r[11],
        "sender":{"id":r[2],"name":r[12],"email":r[13],"role":r[14]},
        "is_mine":r[2]==current_user_id,"reactions":reactions.get(("p",r[0]),{})
    }


def _team_message_payload(r, current_user_id, reactions):
    return {
        "id":r[0],"room_id":r[1],"sender_id":r[2],"body":r[3],"created_at":str(r[4]),
        "edited_at":str(r[5]) if r[5] else None,"deleted_at":str(r[6]) if r[6] else None,
        "is_deleted":bool(r[6]),"reply_to_id":r[7],"reply_body":r[8],"reply_sender_id":r[9],
        "reply_sender_name":r[10],"sender":{"id":r[2],"name":r[11],"email":r[12],"role":r[13]},
        "is_mine":r[2]==current_user_id,"reactions":reactions.get(("t",r[0]),{})
    }


@app.post("/api/chat/presence")
def chat_presence(current_user: dict = Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("""INSERT INTO chat_presence(user_id,last_seen) VALUES(%s,NOW())
                           ON CONFLICT(user_id) DO UPDATE SET last_seen=NOW()""",(current_user["id"],))
        conn.commit()
    return {"success":True,"online":True}


@app.get("/api/chat/users")
def chat_users(current_user: dict = Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("""SELECT u.id,u.name,u.email,COALESCE(u.role,'Agent'),COALESCE(u.active,TRUE),COALESCE(u.approval_status,'Approved'),
                                  p.last_seen
                           FROM users u LEFT JOIN chat_presence p ON p.user_id=u.id
                           WHERE u.id<>%s AND COALESCE(u.active,TRUE)=TRUE AND COALESCE(u.approval_status,'Approved')='Approved'
                           ORDER BY CASE WHEN COALESCE(u.role,'Agent')='Admin' THEN 0 ELSE 1 END,LOWER(COALESCE(u.name,u.email,''))""",(current_user["id"],))
            rows=cur.fetchall()
    return {"users":[{"id":r[0],"name":r[1],"email":r[2],"role":r[3],"active":bool(r[4]),"approval_status":r[5],
                       "online":bool(r[6] and (datetime.now(timezone.utc)-r[6]).total_seconds()<=CHAT_ONLINE_SECONDS),"last_seen":str(r[6]) if r[6] else None} for r in rows]}


@app.get("/api/chat/team-room")
def get_team_room(current_user: dict = Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("SELECT id,room_name FROM chat_team_room WHERE id=1")
            room=cur.fetchone()
            if not room:
                cur.execute("INSERT INTO chat_team_room(id,room_name) VALUES(1,'Team Room') RETURNING id,room_name")
                room=cur.fetchone()
            cur.execute("""SELECT COUNT(*) FROM chat_team_messages m WHERE m.room_id=1 AND m.sender_id<>%s
                           AND m.id>COALESCE((SELECT last_read_message_id FROM chat_team_reads WHERE room_id=1 AND user_id=%s),0)""",(current_user["id"],current_user["id"]))
            unread=int(cur.fetchone()[0] or 0)
        conn.commit()
    return {"room":{"id":room[0],"name":room[1],"unread_count":unread}}


@app.get("/api/chat/team-room/messages")
def get_team_room_messages(limit:int=Query(100,ge=1,le=200),before_id:Optional[int]=Query(None,gt=0),current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            sql="""SELECT m.id,m.room_id,m.sender_id,m.body,m.created_at,m.edited_at,m.deleted_at,m.reply_to_id,
                            rm.body,rm.sender_id,ru.name,u.name,u.email,u.role
                     FROM chat_team_messages m JOIN users u ON u.id=m.sender_id
                     LEFT JOIN chat_team_messages rm ON rm.id=m.reply_to_id
                     LEFT JOIN users ru ON ru.id=rm.sender_id
                     WHERE m.room_id=1"""
            params=[]
            if before_id is not None: sql+=" AND m.id<%s";params.append(before_id)
            sql+=" ORDER BY m.created_at DESC,m.id DESC LIMIT %s";params.append(limit)
            cur.execute(sql,params);rows=list(reversed(cur.fetchall()))
            ids=[r[0] for r in rows];reactions=_chat_reactions(cur,team_ids=ids)
    return {"room":{"id":1,"name":"Team Room"},"messages":[_team_message_payload(r,current_user["id"],reactions) for r in rows]}


@app.post("/api/chat/team-room/messages")
def send_team_room_message(req:ChatMessageRequest,current_user:dict=Depends(get_current_user)):
    body=req.body.strip()
    if not body: raise HTTPException(status_code=400,detail="Message cannot be empty.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            if req.reply_to_id:
                cur.execute("SELECT id FROM chat_team_messages WHERE id=%s AND room_id=1",(req.reply_to_id,))
                if not cur.fetchone(): raise HTTPException(status_code=400,detail="Reply target not found.")
            cur.execute("""INSERT INTO chat_team_messages(room_id,sender_id,body,reply_to_id) VALUES(1,%s,%s,%s) RETURNING id,created_at""",(current_user["id"],body,req.reply_to_id))
            mid,created=cur.fetchone()
            cur.execute("SELECT id FROM users WHERE id<>%s AND COALESCE(active,TRUE)=TRUE AND COALESCE(approval_status,'Approved')='Approved'",(current_user["id"],))
            for (rid,) in cur.fetchall():
                cur.execute("""INSERT INTO notifications(recipient_id,actor_id,lead_id,notification_type,title,message) VALUES(%s,%s,NULL,'chat_message',%s,%s)""",(rid,current_user["id"],f"New team message from {(current_user.get('name') or current_user.get('email') or 'User')}",body[:250]))
        conn.commit()
    return {"success":True,"message":{"id":mid,"room_id":1,"sender_id":current_user["id"],"body":body,"created_at":str(created)}}


@app.patch("/api/chat/team-room/messages/{message_id}")
def edit_team_message(message_id:int,req:ChatEditMessageRequest,current_user:dict=Depends(get_current_user)):
    body=req.body.strip()
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("UPDATE chat_team_messages SET body=%s,edited_at=NOW() WHERE id=%s AND sender_id=%s AND deleted_at IS NULL RETURNING id",(body,message_id,current_user["id"]))
            if not cur.fetchone(): raise HTTPException(status_code=404,detail="Message not found or cannot be edited.")
        conn.commit()
    return {"success":True}


@app.delete("/api/chat/team-room/messages/{message_id}")
def delete_team_message(message_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("UPDATE chat_team_messages SET deleted_at=NOW(),body='' WHERE id=%s AND sender_id=%s AND deleted_at IS NULL RETURNING id",(message_id,current_user["id"]))
            if not cur.fetchone(): raise HTTPException(status_code=404,detail="Message not found or cannot be deleted.")
        conn.commit()
    return {"success":True}


@app.post("/api/chat/team-room/messages/{message_id}/reactions")
def react_team_message(message_id:int,req:ChatReactionRequest,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("SELECT id FROM chat_team_messages WHERE id=%s AND room_id=1 AND deleted_at IS NULL",(message_id,))
            if not cur.fetchone(): raise HTTPException(status_code=404,detail="Message not found.")
            cur.execute("DELETE FROM chat_message_reactions WHERE team_message_id=%s AND user_id=%s AND reaction=%s RETURNING id",(message_id,current_user["id"],req.reaction))
            if not cur.fetchone():
                cur.execute("INSERT INTO chat_message_reactions(team_message_id,user_id,reaction) VALUES(%s,%s,%s)",(message_id,current_user["id"],req.reaction))
        conn.commit()
    return {"success":True}


@app.post("/api/chat/team-room/read")
def mark_team_room_read(current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("SELECT COALESCE(MAX(id),0) FROM chat_team_messages WHERE room_id=1");last_id=int(cur.fetchone()[0] or 0)
            cur.execute("""INSERT INTO chat_team_reads(room_id,user_id,last_read_message_id,last_read_at) VALUES(1,%s,%s,NOW())
                           ON CONFLICT(room_id,user_id) DO UPDATE SET last_read_message_id=EXCLUDED.last_read_message_id,last_read_at=NOW()""",(current_user["id"],last_id))
        conn.commit()
    return {"success":True,"last_read_message_id":last_id}


@app.get("/api/chat/team-room/unread-count")
def team_room_unread_count(current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("""SELECT COUNT(*) FROM chat_team_messages m WHERE m.room_id=1 AND m.sender_id<>%s
                           AND m.id>COALESCE((SELECT last_read_message_id FROM chat_team_reads WHERE room_id=1 AND user_id=%s),0)""",(current_user["id"],current_user["id"]))
            count=int(cur.fetchone()[0] or 0)
    return {"unread_count":count}


@app.get("/api/chat/conversations")
def chat_conversations(current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("""SELECT c.id,c.user1_id,c.user2_id,c.created_at,c.last_message_at,
                                  u1.name,u1.email,u1.role,u2.name,u2.email,u2.role,
                                  COALESCE((SELECT COUNT(*) FROM chat_messages m WHERE m.conversation_id=c.id AND m.sender_id<>%s AND m.is_read=FALSE),0),
                                  (SELECT m.body FROM chat_messages m WHERE m.conversation_id=c.id ORDER BY m.created_at DESC,m.id DESC LIMIT 1),
                                  (SELECT m.sender_id FROM chat_messages m WHERE m.conversation_id=c.id ORDER BY m.created_at DESC,m.id DESC LIMIT 1),
                                  p1.last_seen,p2.last_seen
                           FROM chat_conversations c JOIN users u1 ON u1.id=c.user1_id JOIN users u2 ON u2.id=c.user2_id
                           LEFT JOIN chat_presence p1 ON p1.user_id=u1.id LEFT JOIN chat_presence p2 ON p2.user_id=u2.id
                           WHERE c.user1_id=%s OR c.user2_id=%s ORDER BY c.last_message_at DESC,c.id DESC""",(current_user["id"],current_user["id"],current_user["id"]))
            rows=cur.fetchall()
    out=[]
    for r in rows:
        other_id=r[2] if r[1]==current_user["id"] else r[1];name=r[8] if r[1]==current_user["id"] else r[5];email=r[9] if r[1]==current_user["id"] else r[6];role=r[10] if r[1]==current_user["id"] else r[7];last_seen=r[15] if r[1]==current_user["id"] else r[14]
        out.append({"id":r[0],"other_user":{"id":other_id,"name":name,"email":email,"role":role,"online":bool(last_seen and (datetime.now(timezone.utc)-last_seen).total_seconds()<=CHAT_ONLINE_SECONDS),"last_seen":str(last_seen) if last_seen else None},"created_at":str(r[3]),"last_message_at":str(r[4]),"unread_count":int(r[11] or 0),"last_message":r[12] or "","last_sender_id":r[13]})
    return {"conversations":out}


@app.post("/api/chat/conversations")
def start_chat(req:ChatStartRequest,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            if req.user_id==current_user["id"]: raise HTTPException(status_code=400,detail="You cannot start a chat with yourself.")
            if not _chat_user_is_available(cur,req.user_id): raise HTTPException(status_code=404,detail="That user is not available for chat.")
            user1_id,user2_id=_chat_pair(current_user["id"],req.user_id)
            cur.execute("INSERT INTO chat_conversations(user1_id,user2_id) VALUES(%s,%s) ON CONFLICT(user1_id,user2_id) DO NOTHING RETURNING id",(user1_id,user2_id));row=cur.fetchone()
            if row: cid=row[0]
            else:
                cur.execute("SELECT id FROM chat_conversations WHERE user1_id=%s AND user2_id=%s",(user1_id,user2_id));row=cur.fetchone();cid=row[0] if row else None
            if not cid: raise HTTPException(status_code=500,detail="Unable to create the chat conversation.")
        conn.commit()
    return {"success":True,"conversation_id":cid}


@app.get("/api/chat/conversations/{conversation_id}/messages")
def chat_messages(conversation_id:int,limit:int=Query(100,ge=1,le=200),before_id:Optional[int]=Query(None,gt=0),current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            c=_chat_conversation_row(cur,conversation_id)
            if not c: raise HTTPException(status_code=404,detail="Conversation not found.")
            if not _chat_is_participant(c,current_user["id"]): raise HTTPException(status_code=403,detail="You do not have access to this conversation.")
            sql="""SELECT m.id,m.conversation_id,m.sender_id,m.body,m.is_read,m.created_at,m.edited_at,m.deleted_at,m.reply_to_id,
                          rm.body,rm.sender_id,ru.name,u.name,u.email,u.role
                   FROM chat_messages m JOIN users u ON u.id=m.sender_id
                   LEFT JOIN chat_messages rm ON rm.id=m.reply_to_id LEFT JOIN users ru ON ru.id=rm.sender_id
                   WHERE m.conversation_id=%s""";params=[conversation_id]
            if before_id is not None: sql+=" AND m.id<%s";params.append(before_id)
            sql+=" ORDER BY m.created_at DESC,m.id DESC LIMIT %s";params.append(limit)
            cur.execute(sql,params);rows=list(reversed(cur.fetchall()));reactions=_chat_reactions(cur,private_ids=[r[0] for r in rows])
    return {"conversation_id":conversation_id,"messages":[_chat_message_payload(r,current_user["id"],reactions) for r in rows]}


@app.post("/api/chat/conversations/{conversation_id}/messages")
def send_chat_message(conversation_id:int,req:ChatMessageRequest,current_user:dict=Depends(get_current_user)):
    body=req.body.strip()
    if not body: raise HTTPException(status_code=400,detail="Message cannot be empty.")
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"]);c=_chat_conversation_row(cur,conversation_id)
            if not c: raise HTTPException(status_code=404,detail="Conversation not found.")
            if not _chat_is_participant(c,current_user["id"]): raise HTTPException(status_code=403,detail="Only conversation participants can send messages.")
            if req.reply_to_id:
                cur.execute("SELECT id FROM chat_messages WHERE id=%s AND conversation_id=%s",(req.reply_to_id,conversation_id))
                if not cur.fetchone(): raise HTTPException(status_code=400,detail="Reply target not found.")
            recipient_id=c[2] if c[1]==current_user["id"] else c[1]
            cur.execute("INSERT INTO chat_messages(conversation_id,sender_id,body,is_read,reply_to_id) VALUES(%s,%s,%s,FALSE,%s) RETURNING id,created_at",(conversation_id,current_user["id"],body,req.reply_to_id));mid,created=cur.fetchone()
            cur.execute("UPDATE chat_conversations SET last_message_at=NOW() WHERE id=%s",(conversation_id,))
            cur.execute("INSERT INTO notifications(recipient_id,actor_id,lead_id,notification_type,title,message) VALUES(%s,%s,NULL,'chat_message',%s,%s)",(recipient_id,current_user["id"],f"New message from {(current_user.get('name') or current_user.get('email') or 'User')}",body[:250]))
        conn.commit()
    return {"success":True,"message":{"id":mid,"conversation_id":conversation_id,"sender_id":current_user["id"],"body":body,"is_read":False,"created_at":str(created)}}


@app.patch("/api/chat/conversations/{conversation_id}/messages/{message_id}")
def edit_chat_message(conversation_id:int,message_id:int,req:ChatEditMessageRequest,current_user:dict=Depends(get_current_user)):
    body=req.body.strip()
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"]);c=_chat_conversation_row(cur,conversation_id)
            if not c or not _chat_is_participant(c,current_user["id"]): raise HTTPException(status_code=403,detail="You do not have access to this conversation.")
            cur.execute("UPDATE chat_messages SET body=%s,edited_at=NOW() WHERE id=%s AND conversation_id=%s AND sender_id=%s AND deleted_at IS NULL RETURNING id",(body,message_id,conversation_id,current_user["id"]))
            if not cur.fetchone(): raise HTTPException(status_code=404,detail="Message not found or cannot be edited.")
        conn.commit()
    return {"success":True}


@app.delete("/api/chat/conversations/{conversation_id}/messages/{message_id}")
def delete_chat_message(conversation_id:int,message_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"]);c=_chat_conversation_row(cur,conversation_id)
            if not c or not _chat_is_participant(c,current_user["id"]): raise HTTPException(status_code=403,detail="You do not have access to this conversation.")
            cur.execute("UPDATE chat_messages SET deleted_at=NOW(),body='' WHERE id=%s AND conversation_id=%s AND sender_id=%s AND deleted_at IS NULL RETURNING id",(message_id,conversation_id,current_user["id"]))
            if not cur.fetchone(): raise HTTPException(status_code=404,detail="Message not found or cannot be deleted.")
        conn.commit()
    return {"success":True}


@app.post("/api/chat/conversations/{conversation_id}/messages/{message_id}/reactions")
def react_chat_message(conversation_id:int,message_id:int,req:ChatReactionRequest,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"]);c=_chat_conversation_row(cur,conversation_id)
            if not c or not _chat_is_participant(c,current_user["id"]): raise HTTPException(status_code=403,detail="You do not have access to this conversation.")
            cur.execute("SELECT id FROM chat_messages WHERE id=%s AND conversation_id=%s AND deleted_at IS NULL",(message_id,conversation_id))
            if not cur.fetchone(): raise HTTPException(status_code=404,detail="Message not found.")
            cur.execute("DELETE FROM chat_message_reactions WHERE private_message_id=%s AND user_id=%s AND reaction=%s",(message_id,current_user["id"],req.reaction))
            # Do not use RETURNING id here: older installations may have a
            # legacy reactions table without an id column.
            if cur.rowcount == 0:
                cur.execute("INSERT INTO chat_message_reactions(private_message_id,user_id,reaction) VALUES(%s,%s,%s)",(message_id,current_user["id"],req.reaction))
        conn.commit()
    return {"success":True}


@app.post("/api/chat/conversations/{conversation_id}/read")
def mark_chat_read(conversation_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"]);c=_chat_conversation_row(cur,conversation_id)
            if not c: raise HTTPException(status_code=404,detail="Conversation not found.")
            if not _chat_is_participant(c,current_user["id"]): raise HTTPException(status_code=403,detail="Only conversation participants can mark messages as read.")
            cur.execute("UPDATE chat_messages SET is_read=TRUE WHERE conversation_id=%s AND sender_id<>%s AND is_read=FALSE",(conversation_id,current_user["id"]));count=cur.rowcount
        conn.commit()
    return {"success":True,"marked_read":count}


@app.get("/api/chat/unread-count")
def chat_unread_count(current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            _require_chat_user(cur,current_user["id"])
            cur.execute("SELECT COUNT(*) FROM chat_messages m JOIN chat_conversations c ON c.id=m.conversation_id WHERE (c.user1_id=%s OR c.user2_id=%s) AND m.sender_id<>%s AND m.is_read=FALSE",(current_user["id"],current_user["id"],current_user["id"]));count=cur.fetchone()[0]
    return {"unread_count":int(count)}


# Groups

# ============================================================

def assert_group_owner(cur, group_id: int, user_id: int):

    cur.execute("SELECT id FROM lead_groups WHERE id=%s AND user_id=%s", (group_id, user_id))

    if not cur.fetchone(): raise HTTPException(status_code=404, detail="Group not found.")

@app.get("/api/groups")

def list_groups(current_user: dict = Depends(get_current_user)):

    with db_conn() as conn:

        with conn.cursor() as cur:

            cur.execute("SELECT id,name,created_at FROM lead_groups WHERE user_id=%s ORDER BY created_at DESC", (current_user["id"],))

            rows = cur.fetchall()

    return {"groups": [{"id": r[0], "name": r[1], "created_at": str(r[2])} for r in rows]}

@app.post("/api/groups")

def create_group(req: GroupCreateRequest, current_user: dict = Depends(get_current_user)):

    name = req.name.strip()

    if not name: raise HTTPException(status_code=400, detail="Group name is required.")

    with db_conn() as conn:

        with conn.cursor() as cur:

            cur.execute("INSERT INTO lead_groups(user_id,name) VALUES(%s,%s) RETURNING id,name,created_at", (current_user["id"], name[:150]))

            row = cur.fetchone()
            log_activity(cur,current_user["id"],None,"group_created",{"group_id":row[0],"name":row[1]})

        conn.commit()

    return {"group": {"id": row[0], "name": row[1], "created_at": str(row[2])}}

@app.post("/api/groups/{group_id}/leads")
def add_lead_to_group(group_id: int, req: GroupBulkAddRequest, current_user: dict = Depends(get_current_user)):
    if not req.lead_ids: raise HTTPException(status_code=400, detail="No lead selected.")
    added = 0
    with db_conn() as conn:
        with conn.cursor() as cur:
            assert_group_owner(cur, group_id, current_user["id"])
            for lead_id in req.lead_ids[:500]:
                lead_scope, lead_params = lead_access_clause(current_user)
                cur.execute(f"SELECT id FROM leads WHERE id=%s AND {lead_scope}", (lead_id, *lead_params))
                if cur.fetchone():
                    cur.execute("INSERT INTO group_leads(group_id,lead_id) VALUES(%s,%s) ON CONFLICT DO NOTHING", (group_id, lead_id))
                    added += cur.rowcount
            log_activity(cur,current_user["id"],None,"group_leads_added",{"group_id":group_id,"count":added})
        conn.commit()
    return {"success": True, "added": added}

@app.post("/api/groups/{group_id}/leads/bulk")

def add_leads_to_group_bulk(group_id: int, req: GroupBulkAddRequest, current_user: dict = Depends(get_current_user)):

    if not req.lead_ids: raise HTTPException(status_code=400, detail="No leads selected.")

    added = 0

    with db_conn() as conn:

        with conn.cursor() as cur:

            assert_group_owner(cur, group_id, current_user["id"])

            for lead_id in req.lead_ids[:500]:

                lead_scope, lead_params = lead_access_clause(current_user)
                cur.execute(f"SELECT id FROM leads WHERE id=%s AND {lead_scope}", (lead_id, *lead_params))

                if cur.fetchone():

                    cur.execute("INSERT INTO group_leads(group_id,lead_id) VALUES(%s,%s) ON CONFLICT DO NOTHING", (group_id, lead_id))

                    added += cur.rowcount

            log_activity(cur,current_user["id"],None,"group_leads_added",{"group_id":group_id,"count":added})

        conn.commit()

    return {"success": True, "added": added}

@app.delete("/api/groups/{group_id}/leads/{lead_id}")
def remove_lead_from_group(group_id: int, lead_id: int, current_user: dict = Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            assert_group_owner(cur, group_id, current_user["id"])
            cur.execute("DELETE FROM group_leads WHERE group_id=%s AND lead_id=%s", (group_id, lead_id))
            removed = cur.rowcount
            if removed: log_activity(cur,current_user["id"],lead_id,"group_lead_removed",{"group_id":group_id})
        conn.commit()
    return {"success": True, "removed": removed}

@app.get("/api/groups/{group_id}/export-excel")
def export_group_excel(group_id:int,current_user:dict=Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            assert_group_owner(cur,group_id,current_user["id"])
            scope,params=lead_access_clause(current_user,"l")
            cur.execute("""SELECT l.name,l.phone,l.email,l.address,l.city,l.region,l.country,l.website,l.instagram,l.facebook,l.twitter,l.rating,l.reviews,l.category,l.lead_score,l.lead_type,l.ai_recommendation,l.ai_reason,l.status,l.assigned_to,l.last_contacted_date,l.next_followup_date,l.source,l.maps_url,l.notes,c.call_outcome,c.notes FROM leads l JOIN group_leads gl ON gl.lead_id=l.id LEFT JOIN LATERAL (SELECT call_outcome,notes FROM call_logs WHERE lead_id=l.id ORDER BY call_date DESC,id DESC LIMIT 1)c ON TRUE WHERE gl.group_id=%s AND """+scope+" ORDER BY l.created_at DESC",(group_id,*params))
            rows=cur.fetchall()
    with db_conn() as conn:
        with conn.cursor() as cur:
            log_activity(cur,current_user["id"],None,"group_exported",{"group_id":group_id})
        conn.commit()
    wb=Workbook(); ws=wb.active; ws.title="Group Leads"; ws.append(EXPORT_HEADERS)
    for row in rows: ws.append([excel_safe_value(v) for v in row])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        width=min(max(max(len(str(c.value or "")) for c in col)+2,12),45); ws.column_dimensions[col[0].column_letter].width=width
    buf=BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="group_{group_id}.xlsx"'})


@app.get("/api/groups/{group_id}/export-excel/")
def export_group_excel_trailing(group_id:int,current_user:dict=Depends(get_current_user)):
    return export_group_excel(group_id, current_user)

@app.get("/api/groups/{group_id}/leads")

def get_group_leads(group_id: int, current_user: dict = Depends(get_current_user)):

    with db_conn() as conn:

        with conn.cursor() as cur:

            assert_group_owner(cur, group_id, current_user["id"])

            lead_scope, lead_params = lead_access_clause(current_user,"l")
            cur.execute("""SELECT l.id,l.name,l.address,l.city,l.region,l.country,l.phone,l.website,l.email,l.instagram,l.facebook,l.twitter,l.category,l.rating,l.reviews,l.maps_url,l.source,l.lead_score,l.lead_type,l.ai_recommendation,l.ai_reason,l.status,l.place_id,l.created_at,l.updated_at,l.notes,l.latitude,l.longitude,l.provider_place_id,l.sources,l.enrichment_source,l.enrichment_at,l.assigned_to,l.last_contacted_date,l.next_followup_date,l.is_duplicate,l.duplicate_of FROM leads l JOIN group_leads gl ON gl.lead_id=l.id WHERE gl.group_id=%s AND """+lead_scope+" ORDER BY l.created_at DESC", (group_id, *lead_params))

            rows = cur.fetchall()

    return {"leads": [row_to_lead(r) for r in rows]}

# ============================================================

class AIResearchRequest(BaseModel):
    query: str = Field(min_length=2,max_length=200)
    city: Optional[str] = Field(None,max_length=120)
    lead_id: Optional[int] = None

def _ai_category_context(query: str, lead: Optional[dict]) -> str:
    text = " ".join([
        str(query or ""),
        str((lead or {}).get("name") or ""),
        str((lead or {}).get("category") or ""),
    ]).casefold()

    playbooks = [
        (("salon","spa","beauty","barber","parlour","parlor"), {
            "category": "Salon / Beauty",
            "needs": [
                "new appointment generation and local discovery",
                "conversion of enquiries into booked appointments",
                "repeat visits, packages and customer retention",
                "fast handling of calls, messages and appointment enquiries"
            ],
            "opportunities": [
                "local visibility and review/reputation journey",
                "booking and enquiry follow-up",
                "re-engagement of previous customers",
                "clear presentation of services, pricing and offers"
            ],
            "questions": [
                "Where do most new appointment enquiries come from today?",
                "How quickly does the team respond to missed calls or online enquiries?",
                "How do you bring previous customers back for their next visit?"
            ]
        }),
        (("restaurant","cafe","coffee shop","bakery","food"), {
            "category": "Restaurant / Food",
            "needs": [
                "local discovery and customer acquisition",
                "conversion from online interest to visits or orders",
                "repeat customers and retention",
                "handling enquiries, reservations or order questions"
            ],
            "opportunities": [
                "Google/local presence and reputation",
                "reservation/order enquiry journey",
                "repeat-visit and loyalty opportunities",
                "menu, offer and location information clarity"
            ],
            "questions": [
                "How do customers usually discover you for the first time?",
                "What happens when someone enquires but does not visit or order?",
                "How do you encourage satisfied customers to return?"
            ]
        }),
        (("dentist","dental","clinic","hospital","doctor","medical"), {
            "category": "Healthcare",
            "needs": [
                "patient discovery and appointment enquiries",
                "conversion from enquiry to confirmed appointment",
                "appointment reminders and follow-up",
                "clear trust, service and location information"
            ],
            "opportunities": [
                "local discovery and reputation",
                "appointment enquiry handling",
                "follow-up for unconfirmed enquiries",
                "clear patient information before contact"
            ],
            "questions": [
                "How do new patients normally find the practice?",
                "How are appointment enquiries followed up if they do not book immediately?",
                "Which types of appointments or services are you trying to grow?"
            ]
        }),
        (("gym","fitness","yoga","sports club","wellness"), {
            "category": "Fitness / Wellness",
            "needs": [
                "membership enquiries and trial conversion",
                "local discovery and lead follow-up",
                "member retention and reactivation",
                "clear presentation of plans, classes and schedules"
            ],
            "opportunities": [
                "trial-to-membership conversion",
                "follow-up of enquiries that go cold",
                "member reactivation",
                "local visibility and reviews"
            ],
            "questions": [
                "How do you currently convert trial or membership enquiries?",
                "How quickly are new enquiries followed up?",
                "What usually causes members to stop or become inactive?"
            ]
        }),
        (("real estate","real estate agency","property","realtor","broker"), {
            "category": "Real Estate",
            "needs": [
                "qualified enquiry generation",
                "fast follow-up on property enquiries",
                "lead qualification and appointment conversion",
                "consistent follow-up across longer sales cycles"
            ],
            "opportunities": [
                "property discovery and enquiry capture",
                "lead response speed",
                "follow-up sequences for interested prospects",
                "qualification before agent time is spent"
            ],
            "questions": [
                "Where do most property enquiries originate?",
                "How quickly does an agent respond to a new enquiry?",
                "How are interested prospects followed up after the first conversation?"
            ]
        }),
        (("automobile","automotive","car dealer","car dealership","vehicle dealer","motor"), {
            "category": "Automotive",
            "needs": [
                "vehicle discovery and qualified enquiries",
                "test-drive or showroom appointment conversion",
                "follow-up across longer purchase decisions",
                "service and repeat-customer retention"
            ],
            "opportunities": [
                "vehicle discovery and enquiry journey",
                "test-drive lead follow-up",
                "quotation follow-up",
                "service reminders and repeat engagement"
            ],
            "questions": [
                "How are vehicle enquiries captured and followed up today?",
                "What happens to customers who ask for a quotation but do not buy immediately?",
                "How do you bring existing customers back for service or future purchases?"
            ]
        }),
        (("retail","shop","store","boutique","fashion","clothing","jewellery","jewelry"), {
            "category": "Retail",
            "needs": [
                "local discovery and footfall",
                "conversion of enquiries into purchases",
                "repeat purchases and customer retention",
                "clear product and offer communication"
            ],
            "opportunities": [
                "local visibility and reputation",
                "customer enquiry handling",
                "repeat-customer engagement",
                "product/offer discovery journey"
            ],
            "questions": [
                "How do customers usually discover the store?",
                "How do you follow up with customers who show interest but do not purchase?",
                "How do you encourage repeat purchases?"
            ]
        }),
    ]

    for keywords, data in playbooks:
        if any(k in text for k in keywords):
            return json.dumps(data, ensure_ascii=False)

    generic = {
        "category": str((lead or {}).get("category") or query or "Unknown business category"),
        "needs": [
            "customer discovery and acquisition",
            "conversion of enquiries into customers",
            "follow-up and customer retention",
            "a clear and low-friction customer journey"
        ],
        "opportunities": [
            "local/online visibility",
            "lead and enquiry response",
            "conversion points in the customer journey",
            "repeat-customer engagement"
        ],
        "questions": [
            "How do customers usually discover the business?",
            "What happens after a new enquiry is received?",
            "How are previous customers encouraged to return?"
        ]
    }
    return json.dumps(generic, ensure_ascii=False)

def _ai_research_prompt(query,city,lead,snippets):
    category_context = _ai_category_context(query, lead)
    return f"""
You are a B2B business-intelligence researcher inside a lead-generation CRM.
Research ONE SPECIFIC BUSINESS, not the business category in general.

EVIDENCE PRIORITY
1. Lead Data = CRM facts.
2. Structured public business evidence = strongest external evidence.
3. Public search snippets = supporting evidence.
4. Category Context = reasoning only, never a fact about this business.

RULES
- what_they_have contains ONLY concrete facts supported by Lead Data or public evidence.
- Prefer address, phone, website, hours, services/products, booking/order channels, locations, ratings/review counts, and stated offers.
- Never say the business lacks something merely because it was not found.
- Do not turn review opinions into objective facts.
- what_they_may_need contains 2-4 hypotheses/opportunities, each tied to a verified fact or evidence gap and briefly explaining why.
- Use may/could/worth investigating. Never present a category assumption as a confirmed problem.
- what_we_can_offer maps 1:1 to opportunities and describes solution TYPES only. Never invent the user's company services, pricing, tools, clients, or capabilities.
- how_to_approach identifies a likely decision-maker, uses one verified fact, tests one hypothesis, and asks a discovery question.
- opening_pitch uses a real business fact and a discovery question.
- next_action is one concrete sales action.

BUSINESS QUERY:
{query}
CITY:
{city or "Unknown"}
LEAD DATA:
{json.dumps(lead or {},default=str)[:14000]}
PUBLIC BUSINESS EVIDENCE:
{json.dumps(snippets,default=str)[:24000]}
CATEGORY CONTEXT:
{category_context}

Return ONLY valid JSON with exactly these keys:
{{
  "business_snapshot": "One concise sentence using verified facts only.",
  "what_they_have": [{{"fact":"Concrete verified fact","source":"Exact source/domain or Lead Data"}}],
  "what_they_may_need": ["Potential opportunity + why it is relevant to this business"],
  "what_we_can_offer": ["Potential solution area mapped to the opportunity"],
  "how_to_approach": "2-4 concise sentences",
  "opening_pitch": "2-3 natural sentences",
  "next_action": "One concrete next action"
}}
Keep it specific enough that a salesperson can act in 30 seconds.
"""

def _ai_clean_text(value):
    if value is None: return ""
    return re.sub(r"\s+", " ", str(value)).strip()

def _ai_add_evidence(evidence, fact, source):
    fact=_ai_clean_text(fact); source=_ai_clean_text(source) or "Public search evidence"
    if not fact: return
    if any(x["fact"].casefold()==fact.casefold() for x in evidence): return
    evidence.append({"fact":fact,"source":source})

class _AIPageParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.title=[]; self.meta={}; self.text=[]; self._in_title=False
        self.jsonld=[]; self._in_script=False; self._script_type=""; self._script_buf=[]
    def handle_starttag(self, tag, attrs):
        a=dict(attrs)
        if tag.lower()=="title": self._in_title=True
        if tag.lower()=="meta":
            key=a.get("name") or a.get("property")
            val=a.get("content")
            if key and val: self.meta[key.lower()]=val.strip()
        if tag.lower()=="script" and (a.get("type") or "").lower()=="application/ld+json":
            self._in_script=True; self._script_buf=[]
    def handle_endtag(self, tag):
        if tag.lower()=="title": self._in_title=False
        if tag.lower()=="script" and self._in_script:
            self._in_script=False
            if self._script_buf: self.jsonld.append("".join(self._script_buf))
    def handle_data(self, data):
        if self._in_title: self.title.append(data)
        if self._in_script: self._script_buf.append(data)
        if data.strip(): self.text.append(data.strip())

def _ai_safe_public_url(url):
    try:
        from urllib.parse import urlparse
        import ipaddress
        p=urlparse(url)
        if p.scheme not in ("http","https") or not p.hostname: return False
        host=p.hostname.lower()
        if host in {"localhost","127.0.0.1","0.0.0.0","::1"}: return False
        try:
            ip=ipaddress.ip_address(host)
            if ip.is_private or ip.is_loopback or ip.is_link_local: return False
        except ValueError: pass
        return True
    except Exception:
        return False

def _ai_extract_page(url):
    if not _ai_safe_public_url(url): return []
    evidence=[]
    try:
        r=requests.get(url,headers={"User-Agent":"Mozilla/5.0 (compatible; LeadSearchEngine/1.0)"},timeout=8,allow_redirects=True)
        if r.status_code!=200 or "text/html" not in (r.headers.get("content-type") or "").lower(): return []
        parser=_AIPageParser(); parser.feed(r.text[:800000])
        source=url.split("/",3)[2]
        title=_ai_clean_text(" ".join(parser.title))
        desc=_ai_clean_text(parser.meta.get("description") or parser.meta.get("og:description"))
        if title: _ai_add_evidence(evidence,f"Website title: {title}",source)
        if desc: _ai_add_evidence(evidence,f"Website description: {desc[:500]}",source)
        body=" ".join(parser.text)
        patterns=[
            ("Phone",r"(?:\+?\d[\d\s().-]{7,}\d)"),
            ("Email",r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}"),
        ]
        for label,pat in patterns:
            vals=[]
            for m in re.findall(pat,body,re.I):
                v=_ai_clean_text(m)
                if v and v not in vals: vals.append(v)
            for v in vals[:3]: _ai_add_evidence(evidence,f"{label}: {v}",source)
        for raw in parser.jsonld[:10]:
            try:
                obj=json.loads(raw)
            except Exception: continue
            objs=obj if isinstance(obj,list) else [obj]
            for o in objs:
                if not isinstance(o,dict): continue
                typ=str(o.get("@type") or "")
                if not any(x in typ.casefold() for x in ("localbusiness","organization","restaurant","store","beautysalon","cafe","medicalbusiness")): continue
                mapping=[("Business name","name"),("Address","address"),("Phone","telephone"),("Website","url"),("Opening hours","openingHours"),("Business category","servesCuisine"),("Description","description")]
                for label,key in mapping:
                    val=o.get(key)
                    if isinstance(val,dict):
                        parts=[]
                        for k in ("streetAddress","addressLocality","addressRegion","postalCode","addressCountry"):
                            if val.get(k): parts.append(str(val[k]))
                        val=", ".join(parts)
                    if isinstance(val,list): val="; ".join(str(x) for x in val)
                    if val not in (None,"",[],{}): _ai_add_evidence(evidence,f"{label}: {val}",source)
        # Pull a compact, human-readable services/menu description only when the page exposes it.
        service_hits=[]
        for term in ("services","service","menu","treatments","products","appointments","booking","order online"):
            if re.search(r"\b"+re.escape(term)+r"\b",body,re.I): service_hits.append(term)
        if service_hits: _ai_add_evidence(evidence,"Website sections mention: "+", ".join(service_hits[:8]),source)
    except Exception:
        return []
    return evidence[:20]

def _ai_name_match(query,name,address=""):
    q=set(re.findall(r"[a-z0-9]+",query.casefold()))
    n=set(re.findall(r"[a-z0-9]+",f"{name} {address}".casefold()))
    if not q or not n: return False
    overlap=len(q & n)/max(1,min(len(q),len(n)))
    return overlap >= 0.45 or query.casefold() in f"{name} {address}".casefold() or name.casefold() in query.casefold()

def _ai_public_business_evidence(query, city):
    """Collect layered, business-specific public evidence before asking Gemini to reason."""
    evidence=[]; links=[]
    q=_ai_clean_text(query); location=_ai_clean_text(city); search_q=f'"{q}" {location}'.strip()
    if not SERPAPI_API_KEY:
        return evidence
    # Layer 1: Google organic results. Run a few focused queries to avoid relying on one generic snippet.
    search_queries=[search_q, f'"{q}" {location} official website', f'"{q}" {location} phone hours services']
    seen_links=set()
    for sq in search_queries:
        try:
            data=serpapi_get({"engine":"google","q":sq,"hl":"en","gl":"in"},timeout=15)
            for x in (data.get("organic_results") or [])[:8]:
                title=_ai_clean_text(x.get("title")) or "Google search result"
                snippet=_ai_clean_text(x.get("snippet")); link=_ai_clean_text(x.get("link"))
                if snippet: _ai_add_evidence(evidence,snippet[:600],title+(f" · {link}" if link else ""))
                if link and link not in seen_links and _ai_safe_public_url(link):
                    seen_links.add(link); links.append(link)
        except Exception: pass
    # Layer 2: structured Google Maps result.
    matched_maps=[]
    try:
        maps=serpapi_get({"engine":"google_maps","q":search_q,"type":"search","hl":"en","gl":"in"},timeout=15)
        for item in (maps.get("local_results") or [])[:8]:
            name=_ai_clean_text(item.get("title")); address=_ai_clean_text(item.get("address"))
            if name and _ai_name_match(q,name,address): matched_maps.append(item)
        for item in matched_maps[:2]:
            source="Google Maps"
            for label,key in [("Business name","title"),("Category","type"),("Address","address"),("Phone","phone"),("Website","website"),("Rating","rating"),("Review count","reviews"),("Opening hours","hours"),("Price level","price"),("Description","description")]:
                value=item.get(key)
                if value not in (None,"",[],{}): _ai_add_evidence(evidence,f"{label}: {value}",source)
            for key in ("service_options","extensions"):
                value=item.get(key)
                if value: _ai_add_evidence(evidence,f"Public listing details: {value}",source)
            website=item.get("website")
            if website and _ai_safe_public_url(str(website)): links.insert(0,str(website))
    except Exception: pass
    # Layer 3: fetch a small number of public result pages, prioritizing likely official websites.
    def link_priority(u):
        low=u.casefold(); score=0
        for host in ("facebook.com","instagram.com","justdial.com","magicpin.in","tripadvisor.","zomato.com"):
            if host in low: score-=2
        for marker in ("official","/contact","/about","/services","/menu"):
            if marker in low: score+=1
        return score
    for link in sorted(links,key=link_priority,reverse=True)[:4]:
        for fact in _ai_extract_page(link): _ai_add_evidence(evidence,fact["fact"],fact["source"])
    return evidence[:40]

def _ai_research_prompt(query,city,lead,snippets):
    category_context=_ai_category_context(query,lead)
    return f"""
You are the research analyst inside a B2B lead-generation CRM.
Your job is to turn public evidence about ONE SPECIFIC BUSINESS into useful sales intelligence.
Do not write a generic industry report.

EVIDENCE HIERARCHY
A. Lead Data = CRM fact.
B. Google Maps / official business website = high-confidence public evidence.
C. Other public listings/search results = supporting evidence; reviews are opinions, not hard facts.
D. Category Context = reasoning only. Never present it as a fact about the business.

WHAT THEY HAVE
- Return 4-8 concrete facts when evidence exists: name, category, location/address, phone, website, hours, services/products, booking/order channels, ratings/review count, locations, stated offers.
- Never write “they do not have X” merely because X was not found.
- Never convert a reviewer's opinion into a business fact.
- Each fact must include its source.

SOCIAL MEDIA REVIEW
- Review only social profile URLs and search evidence provided. Never infer that an account does not exist just because search returned none.
- When profiles are absent from the supplied evidence, state exactly: “Not found in public search” or “Not verified”; do not state the business has no social presence.
- Assess observable public presence qualitatively only. Do not invent follower counts, engagement rates, posting frequency, or metrics.
- If evidence supports a possible gap, include relevant social media management, content planning, profile optimization, marketing, enquiry/lead pathways, or engagement as hypotheses in opportunities and solution areas.

WHAT THEY MAY NEED
Return 2-4 BUSINESS-SPECIFIC opportunities.
Each opportunity MUST be an object with:
- opportunity: concise hypothesis
- evidence: which verified fact(s) make this worth investigating
- why: business-specific reasoning, not generic industry advice
- discovery_question: one question a salesperson can ask
Use “may”, “could”, or “worth investigating”. Never claim an unverified problem.

WHAT WE CAN OFFER
Map each opportunity to a potential solution TYPE only.
Do not invent the user's company services, products, pricing, clients, or capabilities.
Each item should contain: solution_area + linked_opportunity.

HOW TO APPROACH
Identify the likely decision-maker, cite one verified fact, test one hypothesis, and give a discovery question.

OPENING PITCH
2-3 natural sentences. Use a real verified fact and one discovery question. Do not make unsupported claims.

NEXT ACTION
One concrete action that can be completed before/at outreach.

BUSINESS QUERY:
{query}
CITY:
{city or "Unknown"}
LEAD DATA:
{json.dumps(lead or {},default=str)[:16000]}
PUBLIC EVIDENCE:
{json.dumps(snippets,default=str)[:30000]}
CATEGORY CONTEXT:
{category_context}

Return ONLY JSON:
{{
  "business_snapshot":"...",
  "research_confidence":"high|medium|low",
  "what_they_have":[{{"fact":"...","source":"..."}}],
  "what_they_may_need":[{{"opportunity":"...","evidence":"...","why":"...","discovery_question":"..."}}],
  "what_we_can_offer":[{{"solution_area":"...","linked_opportunity":"..."}}],
  "how_to_approach":"...",
  "opening_pitch":"...",
  "next_action":"..."
}}
"""

def _ai_normalize_research(result, query, city, lead, evidence):
    result=result if isinstance(result,dict) else {}
    facts=[]
    for x in result.get("what_they_have") or []:
        if isinstance(x,dict):
            fact=_ai_clean_text(x.get("fact")); source=_ai_clean_text(x.get("source")) or "Public evidence"
        else: fact=_ai_clean_text(x); source="Model output — verify"
        if fact: _ai_add_evidence(facts,fact,source)
    # Always supplement model facts from hard evidence, but never fabricate.
    for x in evidence:
        _ai_add_evidence(facts,x.get("fact"),x.get("source"))
    needs=[]
    raw_needs=result.get("what_they_may_need") or result.get("likely_needs") or []
    for x in raw_needs:
        if isinstance(x,dict):
            item={k:_ai_clean_text(x.get(k)) for k in ("opportunity","evidence","why","discovery_question")}
            if item["opportunity"]: needs.append(item)
        elif _ai_clean_text(x): needs.append({"opportunity":_ai_clean_text(x),"evidence":"Validate against the business evidence.","why":"Potential opportunity; not a confirmed problem.","discovery_question":"How do you currently handle this today?"})
    offers=[]
    raw_offers=result.get("what_we_can_offer") or result.get("recommended_offer") or []
    for x in raw_offers:
        if isinstance(x,dict):
            sa=_ai_clean_text(x.get("solution_area") or x.get("solution") or x.get("offer")); lo=_ai_clean_text(x.get("linked_opportunity"))
        else: sa=_ai_clean_text(x); lo=""
        if sa: offers.append({"solution_area":sa,"linked_opportunity":lo})
    if not result.get("business_snapshot"):
        name=(lead or {}).get("name") or query; cat=(lead or {}).get("category") or "business"
        result["business_snapshot"]=f"{name} — {cat}, based on CRM and public evidence." 
    result["what_they_have"]=facts[:10]
    result["what_they_may_need"]=needs[:4]
    result["what_we_can_offer"]=offers[:4]
    result["research_confidence"]=_ai_clean_text(result.get("research_confidence")) or ("high" if len(facts)>=5 else "medium" if facts else "low")
    result["how_to_approach"]=_ai_clean_text(result.get("how_to_approach") or result.get("contact_strategy") or result.get("recommended_approach"))
    result["opening_pitch"]=_ai_clean_text(result.get("opening_pitch"))
    result["next_action"]=_ai_clean_text(result.get("next_action")) or "Verify the first opportunity with the owner or manager."
    return result

@app.post('/api/ai-research')
def ai_research(req:AIResearchRequest,current_user:dict=Depends(get_current_user)):
    lead=None
    with db_conn() as conn:
        with conn.cursor() as cur:
            if req.lead_id is not None:
                lead_owner_check(cur,req.lead_id,current_user)
                cur.execute(LEAD_SELECT+' WHERE id=%s',(req.lead_id,)); r=cur.fetchone(); lead=row_to_lead(r) if r else None
    evidence=_ai_public_business_evidence(req.query,req.city)
    saved_crm_lead = lead
    social_lookup = discover_public_social_profiles(lead or {"name": req.query, "city": req.city})
    if lead is None:
        lead = {"name": req.query, "city": req.city, **{k:v for k,v in social_lookup.items() if k in ("linkedin", "instagram", "facebook", "twitter", "youtube", "tiktok", "social_evidence", "social_status")}}
    else:
        for social_key in ("linkedin", "instagram", "facebook", "twitter", "youtube", "tiktok"):
            if not lead.get(social_key) and social_lookup.get(social_key):
                lead[social_key] = social_lookup[social_key]
        if social_lookup.get("social_evidence"):
            lead["social_evidence"] = list(lead.get("social_evidence") or []) + social_lookup["social_evidence"]
        lead["social_status"] = social_lookup.get("social_status") or lead.get("social_status")
    if lead:
        social_evidence = lead.get("social_evidence") or []
        for item in social_evidence if isinstance(social_evidence, list) else []:
            if isinstance(item, dict) and item.get("url"):
                _ai_add_evidence(evidence, f"Public social profile ({item.get('platform','social')}): {item.get('url')}", item.get("source") or "Saved SerpApi evidence")
        for platform in ("linkedin", "instagram", "facebook", "twitter", "youtube", "tiktok"):
            url = lead.get(platform)
            if url:
                _ai_add_evidence(evidence, f"CRM saved public social profile ({platform}): {url}", "Lead CRM record")
    result={}
    if GEMINI_API_KEY:
        try:
            url=f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={GEMINI_API_KEY}"
            payload={"contents":[{"parts":[{"text":_ai_research_prompt(req.query,req.city,lead,evidence)}]}],"generationConfig":{"temperature":0.15,"responseMimeType":"application/json"}}
            r=requests.post(url,json=payload,timeout=35)
            if r.status_code==200:
                raw=r.json().get("candidates",[{}])[0].get("content",{}).get("parts",[{}])[0].get("text","").strip()
                raw=re.sub(r'^```(?:json)?\s*|\s*```$','',raw,flags=re.I).strip()
                result=json.loads(raw)
        except Exception:
            result={}
    if not result:
        category=json.loads(_ai_category_context(req.query,lead)).get("category") or (lead or {}).get("category") or "business"
        result={"business_snapshot":f"{(lead or {}).get('name') or req.query} — {category}, based on available evidence.","what_they_have":evidence[:8],"what_they_may_need":[],"what_we_can_offer":[],"how_to_approach":"Contact the owner or manager, lead with a verified fact, and validate one opportunity before proposing a solution.","opening_pitch":f"Hi, I was researching {(lead or {}).get('name') or req.query}. I found a few public details and wanted to understand how you currently handle new customer enquiries. Would you be open to a quick conversation?","next_action":"Verify the first opportunity directly with the owner or manager."}
    result=_ai_normalize_research(result,req.query,req.city,lead,evidence)
    result["social_profiles"] = {k: (lead or {}).get(k) for k in ("linkedin", "instagram", "facebook", "twitter", "youtube", "tiktok") if (lead or {}).get(k)}
    result["social_status"] = (lead or {}).get("social_status") or ("not_verified" if not req.lead_id else "not_checked")
    with db_conn() as conn:
        with conn.cursor() as cur: log_activity(cur,current_user["id"],req.lead_id,"ai_research_completed",{"query":req.query,"city":req.city,"evidence_count":len(evidence),"confidence":result.get("research_confidence")})
        conn.commit()
    return {"success":True,"query":req.query,"city":req.city,"lead":saved_crm_lead,"sources":evidence,"research":result}

# Search history / export / health

# ============================================================

@app.get("/api/search-history")
def search_history(current_user: dict = Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,city,category,keyword,city_wide,result_count,sources_used,created_at FROM search_history WHERE user_id=%s ORDER BY created_at DESC LIMIT 50", (current_user["id"],))
            rows = cur.fetchall()
    return {"history": [{"id": r[0], "city": r[1], "category": r[2], "keyword": r[3], "city_wide": r[4], "result_count": r[5], "sources_used": r[6] or [], "created_at": str(r[7])} for r in rows]}

@app.post("/api/search-history/{history_id}/rerun")
def rerun_search(history_id: int, current_user: dict = Depends(get_current_user)):
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT city,category,keyword,city_wide FROM search_history WHERE id=%s AND user_id=%s", (history_id, current_user["id"]))
            row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Search history item not found.")
    req = SearchRequest(city=row[0], category=row[1], keyword=row[2], city_wide=bool(row[3]), limit=10)
    return search(req, current_user)

def _export_lead_rows(current_user:dict,lead_ids:Optional[str]=None,has_website:Optional[bool]=None,has_phone:Optional[bool]=None,has_email:Optional[bool]=None,has_social:Optional[bool]=None,min_score:Optional[int]=None,lead_type:Optional[str]=None,min_rating:Optional[float]=None,category:Optional[str]=None,city:Optional[str]=None,source:Optional[str]=None,status:Optional[str]=None,assigned_to:Optional[int]=None,date_from:Optional[str]=None,date_to:Optional[str]=None):
    scope,params=lead_access_clause(current_user); clauses=[scope]; params=list(params)
    if lead_ids:
        try: ids=[int(x) for x in lead_ids.split(',') if x.strip()][:500]
        except ValueError: raise HTTPException(status_code=400,detail="Invalid lead selection.")
        clauses.append("l.id=ANY(%s)"); params.append(ids)
    for flag,col in [(has_website,'website'),(has_phone,'phone'),(has_email,'email')]:
        if flag is True: clauses.append(f"NULLIF(TRIM({col}),'') IS NOT NULL")
        elif flag is False: clauses.append(f"NULLIF(TRIM({col}),'') IS NULL")
    if has_social is True: clauses.append("(NULLIF(TRIM(instagram),'') IS NOT NULL OR NULLIF(TRIM(facebook),'') IS NOT NULL OR NULLIF(TRIM(twitter),'') IS NOT NULL)")
    if min_score is not None: clauses.append("COALESCE(lead_score,0)>=%s"); params.append(min_score)
    if lead_type: clauses.append("lead_type=%s"); params.append(lead_type)
    if min_rating is not None: clauses.append("COALESCE(rating,0)>=%s"); params.append(min_rating)
    for val,col in [(category,'category'),(city,'city'),(source,'source'),(status,'status')]:
        if val: clauses.append(f"{col} ILIKE %s"); params.append(f"%{val}%")
    if assigned_to is not None: clauses.append("assigned_to=%s"); params.append(assigned_to)
    start_utc, end_utc = local_date_range_utc(date_from, date_to)
    if start_utc is not None: clauses.append("l.created_at >= %s"); params.append(start_utc)
    if end_utc is not None: clauses.append("l.created_at < %s"); params.append(end_utc)
    sql="SELECT l.name,l.phone,l.email,l.address,l.city,l.region,l.country,l.website,l.instagram,l.facebook,l.twitter,l.rating,l.reviews,l.category,l.lead_score,l.lead_type,l.ai_recommendation,l.ai_reason,l.status,l.assigned_to,l.last_contacted_date,l.next_followup_date,l.source,l.maps_url,l.notes,c.call_outcome,c.notes FROM leads l LEFT JOIN LATERAL (SELECT call_outcome,notes FROM call_logs WHERE lead_id=l.id ORDER BY call_date DESC,id DESC LIMIT 1) c ON TRUE WHERE "+" AND ".join(clauses)+" ORDER BY l.created_at DESC"
    with db_conn() as conn:
        with conn.cursor() as cur: cur.execute(sql,params); return cur.fetchall()

EXPORT_HEADERS=["Business Name","Phone","Email","Address","City","Region","Country","Website","Instagram","Facebook","Twitter/X","Rating","Reviews","Category","Lead Score","Lead Priority","AI Recommendation","AI Reason","Status","Assigned To","Last Contacted Date","Next Follow-up Date","Source","Maps URL","Notes","Latest Call Outcome","Latest Call Notes"]

def export_rows(current_user,**kwargs): return _export_lead_rows(current_user,**kwargs)

@app.get("/api/export-excel")
def export_excel(lead_ids:Optional[str]=None,has_website:Optional[bool]=None,has_phone:Optional[bool]=None,has_email:Optional[bool]=None,has_social:Optional[bool]=None,min_score:Optional[int]=Query(None,ge=0,le=100),lead_type:Optional[str]=None,min_rating:Optional[float]=Query(None,ge=0,le=5),category:Optional[str]=None,city:Optional[str]=None,source:Optional[str]=None,status:Optional[str]=None,assigned_to:Optional[int]=None,date_from:Optional[str]=None,date_to:Optional[str]=None,current_user:dict=Depends(get_current_user)):
    rows=_export_lead_rows(current_user,lead_ids,has_website,has_phone,has_email,has_social,min_score,lead_type,min_rating,category,city,source,status,assigned_to,date_from,date_to)
    with db_conn() as conn:
        with conn.cursor() as cur:
            log_activity(cur,current_user["id"],None,"lead_exported",{"format":"excel" if "export-excel" in __name__ else "export"})
        conn.commit()
    wb=Workbook(); ws=wb.active; ws.title="Leads"; ws.append(EXPORT_HEADERS)
    for row in rows: ws.append([excel_safe_value(v) for v in row])
    buffer=BytesIO(); wb.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=leads_export.xlsx"})

@app.get("/api/export-csv")
def export_csv(lead_ids:Optional[str]=None,has_website:Optional[bool]=None,has_phone:Optional[bool]=None,has_email:Optional[bool]=None,has_social:Optional[bool]=None,min_score:Optional[int]=Query(None,ge=0,le=100),lead_type:Optional[str]=None,min_rating:Optional[float]=Query(None,ge=0,le=5),category:Optional[str]=None,city:Optional[str]=None,source:Optional[str]=None,status:Optional[str]=None,assigned_to:Optional[int]=None,date_from:Optional[str]=None,date_to:Optional[str]=None,current_user:dict=Depends(get_current_user)):
    rows=_export_lead_rows(current_user,lead_ids,has_website,has_phone,has_email,has_social,min_score,lead_type,min_rating,category,city,source,status,assigned_to,date_from,date_to)
    with db_conn() as conn:
        with conn.cursor() as cur:
            log_activity(cur,current_user["id"],None,"lead_exported",{"format":"csv"})
        conn.commit()
    import io; sio=io.StringIO(); writer=csv.writer(sio); writer.writerow(EXPORT_HEADERS)
    for row in rows: writer.writerow(row)
    return StreamingResponse(iter([sio.getvalue().encode('utf-8-sig')]),media_type='text/csv',headers={'Content-Disposition':'attachment; filename=leads_export.csv'})

def _google_service_account():
    raw = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
    if not raw:
        return None
    try:
        if raw.startswith("{"):
            return json.loads(raw)
        path = Path(raw)
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None
    return None

def _google_access_token(service_account: dict) -> str:
    private_key = service_account.get("private_key")
    client_email = service_account.get("client_email")
    token_uri = service_account.get("token_uri") or "https://oauth2.googleapis.com/token"
    if not private_key or not client_email:
        raise ProviderError("Google service-account JSON is missing private_key or client_email.")
    now = int(datetime.now(timezone.utc).timestamp())
    assertion = jwt.encode({
        "iss": client_email,
        "scope": "https://www.googleapis.com/auth/spreadsheets",
        "aud": token_uri,
        "iat": now,
        "exp": now + 3600,
    }, private_key, algorithm="RS256", headers={"kid": service_account.get("private_key_id")} if service_account.get("private_key_id") else None)
    r = requests.post(token_uri, data={"grant_type":"urn:ietf:params:oauth:grant-type:jwt-bearer","assertion":assertion}, timeout=15)
    if r.status_code != 200:
        raise ProviderError("Google authorization failed. Check the service-account credentials.")
    data = r.json()
    if not data.get("access_token"):
        raise ProviderError("Google did not return an access token.")
    return data["access_token"]

@app.get("/api/google-sheets/status")
def google_sheets_status(current_user:dict=Depends(get_current_user)):
    configured=bool(os.getenv('GOOGLE_SHEETS_SPREADSHEET_ID') and _google_service_account())
    return {"enabled":configured,"configured":configured,"priority":"low","mode":"optional"}

@app.post("/api/google-sheets/sync")
def google_sheets_sync(current_user:dict=Depends(get_current_user)):
    require_admin(current_user)
    spreadsheet_id = os.getenv("GOOGLE_SHEETS_SPREADSHEET_ID", "").strip()
    service_account = _google_service_account()
    if not spreadsheet_id or not service_account:
        raise HTTPException(status_code=501, detail="Google Sheets is optional. Configure GOOGLE_SHEETS_SPREADSHEET_ID and GOOGLE_SERVICE_ACCOUNT_JSON first.")
    sheet_range = os.getenv("GOOGLE_SHEETS_RANGE", "Leads!A1")
    try:
        access_token = _google_access_token(service_account)
        rows = _export_lead_rows(current_user)
        values = [EXPORT_HEADERS] + [list(r) for r in rows]
        headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}
        clear_url = f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/{quote_plus(sheet_range)}:clear"
        clear = requests.post(clear_url, headers=headers, json={}, timeout=20)
        if clear.status_code not in (200, 204):
            raise ProviderError("Google Sheets clear operation failed. Check spreadsheet sharing and range.")
        update_url = f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/{quote_plus(sheet_range)}"
        update = requests.put(update_url, headers=headers, params={"valueInputOption":"RAW"}, json={"range":sheet_range,"majorDimension":"ROWS","values":values}, timeout=30)
        if update.status_code != 200:
            raise ProviderError("Google Sheets update failed. Make sure the service account has Editor access to the sheet.")
        return {"success":True,"message":f"Synced {len(rows)} lead(s) to Google Sheets.","rows":len(rows)}
    except ProviderError as exc:
        raise HTTPException(status_code=502, detail=str(exc))
    except requests.RequestException:
        raise HTTPException(status_code=502, detail="Google Sheets is temporarily unavailable.")
    except Exception as exc:
        print("Google Sheets sync error:", type(exc).__name__)
        raise HTTPException(status_code=502, detail="Google Sheets sync failed.")

@app.get("/api/health")

def health():

    db_connected = False

    if DATABASE_URL:

        try:

            with db_conn() as conn:

                with conn.cursor() as cur:

                    cur.execute("SELECT 1"); cur.fetchone()
                    cur.execute("SELECT to_regclass('public.users'), to_regclass('public.leads'), to_regclass('public.call_logs'), to_regclass('public.activity_logs')")
                    tables = cur.fetchone()
                    db_connected = all(tables)

        except Exception:

            pass

    return {"status": "ok" if db_connected else "degraded", "database_configured": bool(DATABASE_URL), "database_connected": db_connected, "latlng_configured": bool(LATLNG_API_KEY), "serpapi_configured": bool(SERPAPI_API_KEY), "gemini_configured": bool(GEMINI_API_KEY), "firebase_configured": bool(FIREBASE_API_KEY), "jwt_configured": bool(JWT_SECRET), "crm_statuses": sorted(VALID_STATUSES), "roles_enabled": True, "followups_enabled": True, "call_logging_enabled": True, "google_sheets_optional": True}